"""Shared parsing helpers for stream uscourts_caseload (U.S. Courts Judicial Business and Judicial Facts and Figures).

Rules kept everywhere:
  * values are copied as published; a printed dash is kept as the marker DASH (shown as a dash, never imputed);
  * a blank or unreadable cell is None ("Not recorded");
  * district/circuit labels are normalised deterministically (case, spacing, footnote markers) and mapped to the
    CourtListener-style court ids used by court_spine; anything unmapped stays unmapped.
"""
import functools
import re
import unicodedata
from pathlib import Path

ROOT = Path(__file__).resolve().parents[3]
ACQ = ROOT / "private" / "acq" / "uscourts_caseload"
DASH = "-"
DASHES = ("-", "--", "—", "–", "−")
SUPERSCRIPTS = "¹²³⁴⁵⁶⁷⁸⁹⁰"
ORDINAL = re.compile(r"^(\d{1,2})(ST|ND|RD|TH)$", re.I)
CIRCUIT_IDS = {"DC": "cadc", **{f"{n}{'st' if n == 1 else 'nd' if n == 2 else 'rd' if n == 3 else 'th'}": f"ca{n}" for n in range(1, 12)}}
CIRCUIT_NAMES = {"cadc": "D.C. Circuit", "ca1": "First Circuit", "ca2": "Second Circuit", "ca3": "Third Circuit", "ca4": "Fourth Circuit",
                 "ca5": "Fifth Circuit", "ca6": "Sixth Circuit", "ca7": "Seventh Circuit", "ca8": "Eighth Circuit", "ca9": "Ninth Circuit",
                 "ca10": "Tenth Circuit", "ca11": "Eleventh Circuit"}


DISTRICT_LABELS = (
    "DC ME MA NH RI PR CT NY,N NY,E NY,S NY,W VT DE NJ PA,E PA,M PA,W VI MD NC,E NC,M NC,W SC VA,E VA,W WV,N WV,S LA,E LA,M LA,W "
    "MS,N MS,S TX,N TX,E TX,S TX,W KY,E KY,W MI,E MI,W OH,N OH,S TN,E TN,M TN,W IL,N IL,C IL,S IN,N IN,S WI,E WI,W AR,E AR,W IA,N IA,S "
    "MN MO,E MO,W NE ND SD AK AZ CA,N CA,E CA,C CA,S HI ID MT NV OR WA,E WA,W GU NMI CO KS NM OK,N OK,E OK,W UT WY AL,N AL,M AL,S "
    "FL,N FL,M FL,S GA,N GA,M GA,S"
).split()
# Labels that appear only in some tables/editions: 'AR' (the two Arkansas bankruptcy courts are one court from FY2018; also
# the misprint for Arizona in Table S-19), 'IL,E' (the historic Eastern District of Illinois, listed with a footnote in
# Table S-19 for its cumulative counts) and 'AK,E'/'AK,W' (FY2016 Table S-19 misprints of AR,E/AR,W).
EXTRA_LABELS = ("AR", "IL,E", "AK,E", "AK,W")


@functools.lru_cache(maxsize=None)
def _letter_keys(extras=()):
    keys = {}
    for label in DISTRICT_LABELS + [e for e in extras if e in EXTRA_LABELS]:
        keys.setdefault(re.sub(r"[^A-Z]", "", label), []).append(label)
    return {k: v[0] for k, v in keys.items() if len(v) == 1}


class _KeyMaps(dict):
    def __missing__(self, extras):
        return _letter_keys(extras)


LETTER_KEYS = _KeyMaps()


def canon_district_label(text, extras=()):
    """Printed district label -> the canonical published label, or None.
    Deterministic normalisation only: case, spaces, a trailing footnote asterisk, a missing comma ('WIW' -> 'WI,W',
    'MS S' -> 'MS,S'), 'WVA' -> 'WV' (an alternative state abbreviation used in FY2007-FY2010 editions) and 'GUAM' -> 'GU'.
    A key shared by two canonical labels is never resolved; labels in EXTRA_LABELS resolve only when the caller lists them
    in `extras` (so a misprinted 'AR' for Arizona in an edition of a table without the merged Arkansas court is not bound)."""
    up = re.sub(r"[\s*]", "", str(text or "").upper())
    up = up.replace("GUAM", "GU")
    if up.startswith("WVA"):
        up = "WV" + up[3:]
    keys = LETTER_KEYS[tuple(sorted(extras))]
    return keys.get(re.sub(r"[^A-Z]", "", up))


class RepairedInt(int):
    """A count whose PDF text layer printed the thousands separator as a period ('14.110'); the repair is accepted only
    when the table's own totals confirm it (see uc_jb.reconcile)."""
    raw = None


def repaired_count(token):
    m = re.fullmatch(r"(\d{1,3})\.(\d{3})", token)
    if not m:
        return None
    value = RepairedInt(int(m.group(1) + m.group(2)))
    value.raw = token
    return value


def norm_text(value):
    """NFKC, Excel carriage-return escapes removed, whitespace collapsed (no other change)."""
    if value is None:
        return ""
    text = unicodedata.normalize("NFKC", str(value).replace("_x000D_", " ").replace("\xa0", " "))
    return re.sub(r"\s+", " ", text).strip()


def split_footnote(label):
    """'1990 ³' -> ('1990', '3'); 'Franchise ¹' -> ('Franchise', '1'); 'IL,E 1' -> ('IL,E', '1'). Raw superscripts only,
    plus a single trailing digit after a district code or a 4-digit year (the source's footnote placement)."""
    raw = str(label or "")
    marks = "".join(ch for ch in raw if ch in SUPERSCRIPTS)
    text = norm_text("".join(ch for ch in raw if ch not in SUPERSCRIPTS))
    for pattern in (r"^((?:19|20)\d\d)\s+(\d)$", r"^([A-Z]{2}, ?[A-Z])\s?(\d)$", r"^([A-Z]{2,4})\s(\d)$"):
        m = re.match(pattern, text)
        if m:
            return m.group(1), (marks + m.group(2)) or None
    if marks:
        marks = unicodedata.normalize("NFKC", marks)
    return text, (marks or None)


def to_value(value):
    """Published cell -> number, DASH, or None. Strings like '50,072' become numbers; '(X)' style text -> None."""
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return int(value) if isinstance(value, float) and value.is_integer() else value
    text = norm_text(value)
    if text in DASHES:
        return DASH
    text = text.replace(",", "").replace("$", "").strip()
    if text.endswith("%"):
        text = text[:-1]
    if re.fullmatch(r"-?\d+(\.\d+)?", text):
        number = float(text)
        return int(number) if "." not in text else number
    return None


def num(v):
    """Numeric value or None (dash and missing both -> None) for arithmetic checks."""
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def zero_if_dash(v):
    """For additive checks only: a printed dash in a count column means none were reported."""
    if v == DASH:
        return 0
    return num(v)


def norm_court_label(label):
    """'NY, N' -> 'NY,N'; '1ST' -> '1st'; 'TOTAL'/'Total' -> 'Total'; 'GUAM' -> 'GU'."""
    text, _ = split_footnote(label)
    text = text.strip()
    up = text.upper().replace(" ", "")
    if up == "TOTAL":
        return "Total"
    m = ORDINAL.match(up)
    if m:
        n = int(m.group(1))
        return f"{n}{'st' if n == 1 else 'nd' if n == 2 else 'rd' if n == 3 else 'th'}"
    if up == "GUAM":
        return "GU"
    if re.fullmatch(r"[A-Z]{2},[A-Z]", up):
        return up
    if re.fullmatch(r"[A-Z]{2,3}", up):
        return up
    return text


def district_id(label):
    """Published district label -> CourtListener district id (court_spine); None if the label is not a district."""
    label = norm_court_label(label)
    if label == "NMI":
        return "nmid"
    if label == "GU":
        return "gud"
    m = re.fullmatch(r"([A-Z]{2}),([A-Z])", label)
    if m:
        return (m.group(1) + m.group(2) + "d").lower()
    if re.fullmatch(r"[A-Z]{2}", label):
        return label.lower() + "d"
    return None


def circuit_id(label):
    return CIRCUIT_IDS.get(norm_court_label(label))


# ---------------------------------------------------------------- xlsx

def xlsx_grids(path):
    """[(sheet_title, grid)] with merged ranges filled from their top-left cell (values only)."""
    import warnings

    import openpyxl

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(path, data_only=True)
    out = []
    for ws in wb.worksheets:
        grid = [list(row) for row in ws.iter_rows(values_only=True)]
        for merged in ws.merged_cells.ranges:
            value = ws.cell(merged.min_row, merged.min_col).value
            for r in range(merged.min_row, merged.max_row + 1):
                for c in range(merged.min_col, merged.max_col + 1):
                    if r - 1 < len(grid) and c - 1 < len(grid[r - 1]):
                        grid[r - 1][c - 1] = value
        out.append((ws.title, grid))
    return out


def xlsx_raw_grids(path):
    """[(sheet_title, grid)] without merged filling (data rows never use merged cells; headers may)."""
    import warnings

    import openpyxl

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(path, data_only=True)
    return [(ws.title, [list(row) for row in ws.iter_rows(values_only=True)]) for ws in wb.worksheets]


def grid_text(grid, rows=8):
    return " ".join(norm_text(c) for row in grid[:rows] for c in row if c is not None)


# ---------------------------------------------------------------- pdf

def char_lines(page, line_tolerance=2.0, gap=1.0):
    """Text lines rebuilt from character positions (same method as scripts/enrichment/court_stats/tables.py):
    characters are clustered into lines by their top coordinate, ordered by x, and a space is inserted only where
    there is a visible gap between printed glyphs, so a printed number comes back as one token."""
    chars = sorted((c for c in page.chars if c.get("text") and not c["text"].isspace()), key=lambda c: (round(c["top"], 1), c["x0"]))
    clusters = []
    for c in chars:
        if clusters and abs(c["top"] - clusters[-1]["top"]) <= line_tolerance:
            clusters[-1]["chars"].append(c)
        else:
            clusters.append({"top": c["top"], "chars": [c]})
    lines = []
    for cluster in clusters:
        parts, last_x1 = [], None
        for c in sorted(cluster["chars"], key=lambda c: c["x0"]):
            if last_x1 is not None and c["x0"] - last_x1 > gap:
                parts.append(" ")
            parts.append(c["text"])
            last_x1 = c["x1"]
        lines.append(re.sub(r"\s+", " ", "".join(parts)).strip())
    return [line for line in lines if line]


def pdf_lines(path, method="chars"):
    """[(page_no, line)] for every page. method 'chars' (glyph positions) or 'text' (pdfplumber extract_text)."""
    import pdfplumber

    out = []
    with pdfplumber.open(path) as pdf:
        for page_no, page in enumerate(pdf.pages):
            lines = char_lines(page) if method == "chars" else (page.extract_text() or "").splitlines()
            out.extend((page_no, norm_text(line)) for line in lines if norm_text(line))
    return out


TOKEN = r"(?:-?[\d,]*\.?\d+|--?|—|–)"
NUMBERS_TAIL = re.compile(rf"((?:\s+{TOKEN})+)\s*$")


def split_tail(line, n=None):
    """Split a text line into (label, [tokens]) where tokens are the trailing numeric/dash tokens.
    With n, exactly the last n tokens are values and anything before is label (footnote digits stay in the label)."""
    line = norm_text(line)
    parts = line.split(" ")
    tokens = []
    while parts and re.fullmatch(TOKEN, parts[-1]):
        tokens.insert(0, parts.pop())
    if n is not None:
        if len(tokens) < n:
            return line, None
        extra = tokens[:-n]
        tokens = tokens[-n:]
        parts = parts + extra
    return " ".join(parts).strip(), [to_value(t) for t in tokens]


def _values(tokens, repair):
    out = []
    for t in tokens:
        fixed = repaired_count(t) if repair else None
        out.append(fixed if fixed is not None else to_value(t))
    return out


def _resolve_court_parts(parts, extras=(), repair=False):
    cut = next((i for i, p in enumerate(parts) if re.fullmatch(TOKEN, p)), len(parts))
    head, rest = parts[:cut], parts[cut:]
    if not head or len(head) > 3 or not rest or not all(re.fullmatch(TOKEN, p) for p in rest):
        return None
    text = "".join(head)
    if len(text) > 7:
        return None
    if text.upper().rstrip("*") == "TOTAL":
        return "Total", _values(rest, repair)
    if ORDINAL.match(text.upper()):
        return norm_court_label(text), _values(rest, repair)
    if text.upper() == "DC":
        return "DC", _values(rest, repair)
    label = canon_district_label(text, extras)
    if label is None:
        return None
    return label, _values(rest, repair)


def court_row(line, extras=(), repair=False):
    """(normalised court label, [tokens]) for a PDF line that starts with a circuit/district/Total label and carries at
    least one value token, else None. Some editions print the page number as rotated glyphs in the margin; those digits
    land at the start of a line ("8 NY, N 9,966 ...") and are dropped when the rest of the line is a court row (a number
    never precedes a court label in these tables). Footnote digits printed after a label stay among the tokens; the
    caller's width check resolves them. Printing variants resolved by canon_district_label: mixed case ('ToTAl', 'gA,N'),
    missing commas ('WIW'), a space ('MS S'), an asterisk ('IL,E*'), 'WVA'. A line made only of a label (a column heading
    such as 'Total') is not a row."""
    parts = norm_text(line).split(" ")
    if len(parts) > 1 and re.fullmatch(r"\d{1,3}", parts[0]):
        dropped = _resolve_court_parts(parts[1:], extras, repair)
        if dropped:
            return dropped
    return _resolve_court_parts(parts, extras, repair)


def pdf_words_lines(page, tolerance=2.5):
    """[(top, [(x0, text), ...])] word lines with positions, for column assignment."""
    words = sorted(page.extract_words(keep_blank_chars=False, use_text_flow=False), key=lambda w: (round(w["top"], 1), w["x0"]))
    lines = []
    for w in words:
        if lines and abs(w["top"] - lines[-1][0]) <= tolerance:
            lines[-1][1].append((w["x0"], w["text"]))
        else:
            lines.append((w["top"], [(w["x0"], w["text"])]))
    return [(top, sorted(ws)) for top, ws in lines]
