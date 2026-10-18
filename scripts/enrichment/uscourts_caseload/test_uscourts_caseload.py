"""Offline unit tests for the caseload stream (synthetic workbooks and payloads only; no network, no private files).

Run: cd scripts/enrichment/uscourts_caseload && python3 -B -m unittest -v test_uscourts_caseload
"""
import os
import tempfile
import unittest

import build_uscourts_caseload as B
import uc_common as U
import uc_jb as J
import uc_jff as F

D = U.DASH
YEAR_2024 = B.YEARS.index(2024)


class CommonTests(unittest.TestCase):
    def test_district_labels_resolve_only_deterministic_variants(self):
        self.assertEqual(U.canon_district_label("WIW"), "WI,W")
        self.assertEqual(U.canon_district_label("MS S"), "MS,S")
        self.assertEqual(U.canon_district_label("GUAM"), "GU")
        self.assertEqual(U.canon_district_label("ToTAl"), None)           # not a district
        self.assertEqual(U.canon_district_label("IL,E"), None)             # historic label resolves only when the caller allows it
        self.assertEqual(U.canon_district_label("IL,E", extras=("IL,E",)), "IL,E")
        self.assertEqual(U.canon_district_label("AR"), None)               # ambiguous (Arizona misprint / merged Arkansas): not bound by default

    def test_district_and_circuit_ids(self):
        self.assertEqual(U.district_id("NY,S"), "nysd")
        self.assertEqual(U.district_id("NMI"), "nmid")
        self.assertEqual(U.district_id("GU"), "gud")
        self.assertEqual(U.district_id("AK"), "akd")
        self.assertEqual(U.circuit_id("2ND"), "ca2")
        self.assertEqual(U.circuit_id("DC"), "cadc")
        self.assertIsNone(U.circuit_id("NY,S"))

    def test_values_footnotes_and_repairs(self):
        self.assertEqual(U.to_value("50,072"), 50072)
        self.assertEqual(U.to_value("10.4%"), 10.4)
        self.assertEqual(U.to_value("-"), D)
        self.assertIsNone(U.to_value("(X)"))
        self.assertIsNone(U.to_value(None))
        self.assertEqual(U.split_footnote("1990 ³"), ("1990", "3"))
        self.assertEqual(U.split_footnote("IL,E 1"), ("IL,E", "1"))
        fixed = U.repaired_count("14.110")
        self.assertEqual(int(fixed), 14110)
        self.assertEqual(fixed.raw, "14.110")
        self.assertIsNone(U.repaired_count("14.11"))

    def test_court_row_drops_rotated_page_digits_and_reads_dashes(self):
        label, values = U.court_row("8 NY, N 9,966 10 -")
        self.assertEqual(label, "NY,N")
        self.assertEqual(values, [9966, 10, D])
        self.assertIsNone(U.court_row("Total"))                           # a heading alone is not a row
        self.assertEqual(U.court_row("ToTAl 5 6")[0], "Total")


class JbParseTests(unittest.TestCase):
    def rows(self, dc_twice=False):
        base = [("Total", [10, 6]), ("DC", [1, 1]), ("1st", [4, 2]), ("ME", [1, 1]), ("MA", [3, 1]), ("2nd", [5, 3]), ("CT", [5, 3])]
        if dc_twice:
            base.insert(1, ("DC", [1, 1]))
        return base

    def test_dc_printed_once_is_circuit_and_district(self):
        entries = J.classify(self.rows())
        dc = next(e for e in entries if e["label"] == "DC")
        self.assertEqual((dc["court_id"], dc["circuit"]), ("dcd", "cadc"))
        self.assertTrue(dc["dc_single"])

    def test_dc_printed_twice_is_circuit_then_district(self):
        entries = J.classify(self.rows(dc_twice=True))
        kinds = [(e["label"], e["kind"]) for e in entries if e["label"] == "DC"]
        self.assertEqual(kinds, [("DC", "circuit"), ("DC", "district")])

    def test_additive_checks_pass_when_districts_sum_to_circuits_and_total(self):
        entries = J.classify(self.rows())
        checks = []
        failures, flags, repaired = J.additive_checks("T", 2025, entries, ["a", "b"], checks, ["a", "b"])
        self.assertEqual(failures, [])
        self.assertTrue(all(c["ok"] for c in checks))

    def test_additive_checks_flag_and_withhold_an_unconfirmed_repair(self):
        rows = self.rows()
        rows[3] = ("ME", [U.repaired_count("14.110"), 1])                 # a period-for-comma misprint inside a group that does not add up
        entries = J.classify(rows)
        checks = []
        failures, flags, repaired = J.additive_checks("T", 2025, entries, ["a", "b"], checks, ["a"])
        self.assertTrue(failures)
        me = next(e for e in entries if e["label"] == "ME")
        self.assertIsNone(me["values"][0])                                # withheld, never guessed
        self.assertTrue(any("NOT confirmed" in r for r in repaired))
        self.assertFalse(all(c["ok"] for c in checks))


def make_workbook(path, rows, merges=()):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    for m in merges:
        ws.merge_cells(m)
    wb.save(path)


class JffTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)

    def book(self, name, data_rows):
        path = os.path.join(self.tmp.name, name)
        rows = [["Table 9.9/U.S. Courts", None, None, None, None],
                ["Fiscal Year", "Filed", None, "Pending", "Pending"],
                [None, "Cases", "Defendants", "Cases", "Defendants"]] + data_rows
        make_workbook(path, rows, merges=["B2:C2", "D3:D3"])
        return path

    def test_read_sheet_columns_rows_dashes_and_cell_addresses(self):
        path = self.book("a.xlsx", [["2025", 100, 120, 7, "-"], ["1990 ²", 50, None, 3, 4], ["Source: test"]])
        sheet = F.read_sheet(path)
        labels = [c["label"] for c in sheet["columns"]]
        self.assertEqual(labels, ["Filed / Cases", "Filed / Defendants", "Pending / Cases", "Pending / Defendants"])
        self.assertEqual(sorted(sheet["rows"]), [1990, 2025])
        self.assertEqual(sheet["rows"][2025]["row"], 4)
        self.assertEqual(sheet["rows"][1990]["mark"], "2")
        self.assertEqual(sheet["columns"][0]["letter"], "B")
        self.assertIn("Source: test", sheet["notes"])

    def test_merge_prefers_latest_edition_keeps_revisions_and_dashes(self):
        old = F.read_sheet(self.book("old.xlsx", [["2024", 90, 100, 5, 6], ["2023", 80, 90, 4, 5]]))
        new = F.read_sheet(self.book("new.xlsx", [["2025", 100, 120, 7, "-"], ["2024", 91, 100, 5, 6]]))
        merged = F.merge_editions([(2024, {"path": "old.xlsx"}, old), (2025, {"path": "new.xlsx"}, new)])
        self.assertEqual(merged["years"], [2023, 2024, 2025])
        filed = merged["values"]["Filed / Cases"]
        self.assertEqual(filed, [80, 91, 100])                              # 2024 taken from the latest edition that lists it
        self.assertEqual(len(merged["revisions"]), 1)
        rev = merged["revisions"][0]
        self.assertEqual((rev["year"], rev["earlier"], rev["later"]), (2024, 90, 91))
        self.assertEqual(merged["dash"], {"Pending / Defendants": [2025]})   # a printed dash is not a number
        self.assertIsNone(merged["values"]["Pending / Defendants"][2])
        self.assertEqual(merged["cells"][2025]["Filed / Cases"], {"edition": 2025, "address": "B4"})
        self.assertEqual(merged["cells"][2023]["Filed / Cases"], {"edition": 2024, "address": "B5"})

    def test_column_missing_in_an_older_edition_is_not_filled(self):
        old_rows = [["Table 9.9", None, None], ["Fiscal Year", "Filed", None], [None, "Cases", None], ["2024", 90, None]]
        path = os.path.join(self.tmp.name, "o.xlsx")
        make_workbook(path, old_rows)
        old = F.read_sheet(path)
        new = F.read_sheet(self.book("n.xlsx", [["2025", 100, 120, 7, 8]]))
        merged = F.merge_editions([(2024, {"path": "o"}, old), (2025, {"path": "n"}, new)])
        self.assertTrue(merged["skipped"])
        self.assertIsNone(merged["values"]["Pending / Cases"][0])           # 2024 pending is Not recorded, never copied across

    def test_repeated_headings_stay_distinct(self):
        rows = [["Table 9.8", None, None], ["Fiscal Year", "Total", "Total"], ["2025", 1, 2]]
        path = os.path.join(self.tmp.name, "dup.xlsx")
        make_workbook(path, rows)
        sheet = F.read_sheet(path)
        merged = F.merge_editions([(2025, {"path": "d"}, sheet)])
        self.assertEqual(len({c["label"] for c in merged["columns"]}), 2)
        self.assertEqual([merged["values"][c["label"]][0] for c in merged["columns"]], [1, 2])

    def test_numbers_printed_with_footnote_markers_are_read_with_their_marker(self):
        self.assertEqual(F.cell_value("302*"), (302, False, "*"))
        self.assertEqual(F.cell_value("730 [4]"), (730, False, "[4]"))
        self.assertEqual(F.cell_value("1,399,496 \u00b2"), (1399496, False, "2"))
        self.assertEqual(F.cell_value("663 6"), (663, False, "6"))
        self.assertEqual(F.cell_value("-"), (None, True, None))
        self.assertEqual(F.cell_value("12 3"), (None, False, None))           # two small numbers: not a value with a marker
        self.assertEqual(F.cell_value("n/a"), (None, False, None))
        self.assertEqual(F.cell_value(41), (41, False, None))

    def test_marked_values_survive_the_merge_and_report_their_marker(self):
        path = os.path.join(self.tmp.name, "mk.xlsx")
        make_workbook(path, [["Table 1.1", None], ["Fiscal Year", "Judges"], ["2002", "302*"], ["2003", 310]])
        merged = F.merge_editions([(2022, {"path": "mk"}, F.read_sheet(path))])
        self.assertEqual(merged["values"]["Judges"], [302, 310])
        self.assertEqual(merged["marks"], {"Judges": {2002: "*"}})

    def test_sheet_without_year_rows_is_not_a_table(self):
        path = os.path.join(self.tmp.name, "none.xlsx")
        make_workbook(path, [["Table 4.2", None], ["Circuit", "Filed"], ["1st", 5]])
        self.assertIsNone(F.read_sheet(path))


def payload(table, year, rows, ok=True):
    return {"table": table, "year": year, "url": f"https://example.test/{table}-{year}.xlsx", "page_url": f"https://example.test/{table}/{year}", "path": f"p/{table}-{year}", "sha256": "0" * 64,
            "retrieved_at": "2026-10-18T00:00:00+00:00", "format": "xlsx", "checks": [{"check": "arith", "ok": ok, "detail": "x"}], "rows": rows}


def drow(court, circuit, fields, unreconciled=()):
    return {"label": court, "kind": "district", "court_id": court, "circuit": circuit, "fields": fields, "unreconciled": list(unreconciled)}


def trow(fields):
    return {"label": "Total", "kind": "total", "court_id": None, "circuit": None, "fields": fields, "unreconciled": []}


class BuilderTests(unittest.TestCase):
    def jb(self):
        jb = B.JB.__new__(B.JB)
        c4a = lambda total, trial: {"total": total, "no_action": 1, "court_total": total - 1, "before_pretrial": total - 3, "during_after_pretrial": 2, "trial_total": trial, "trial_nonjury": trial, "trial_jury": 0, "pct_trial": 0.1}  # noqa: E731
        t1 = lambda n: {"all_trials": n, "civil_total": n, "civil_nonjury": n, "civil_jury": 0, "criminal_total": 0, "criminal_nonjury": 0, "criminal_jury": 0}  # noqa: E731
        s19 = lambda p: {"in_12m": 1, "out_12m": D, "in_cum": 10, "out_cum": 5, "pending": p}  # noqa: E731
        tables = {}
        for year, (a, b) in {2024: (1000, 600), 2025: (1100, 400)}.items():
            tables[("C-4A", year)] = payload("C-4A", year, [trow(c4a(a + b, 20)), drow("med", "ca1", c4a(a, 10)), drow("nhd", "ca1", c4a(b, 10))])
            tables[("T-1", year)] = payload("T-1", year, [trow(t1(7)), drow("med", "ca1", t1(3)), drow("nhd", "ca1", t1(4))])
            tables[("S-19", year)] = payload("S-19", year, [trow(s19(50)), drow("med", "ca1", s19(30), unreconciled=["out_12m"]), drow("nhd", "ca1", s19(20))], ok=(year == 2025))
        jb.tables = tables
        return jb

    def ctx(self):
        spine = {"districts": {"med": ["District Court, D. Maine", "ME"], "nhd": ["District Court, D. New Hampshire", "NH"]}, "bankruptcy": {"meb": ["Bankruptcy Court, D. Maine", "ME"], "arb": ["x", "AZ"], "nebraskab": ["y", "NE"]}}
        return B.Context(self.jb(), spine, {"circuits": {}}, {"ids": ["ME", "NH"]}, [])

    def test_helpers(self):
        self.assertEqual(B.ratio_pct(1, 4), 25.0)
        self.assertIsNone(B.ratio_pct(1, 0))
        self.assertIsNone(B.ratio_pct(None, 4))
        self.assertEqual(B.pct_change(150, 100), 50.0)
        self.assertIsNone(B.pct_change(5, 0))
        self.assertEqual(B.competition_rank({"a": 5, "b": 5, "c": 1}), {"a": 1, "b": 1, "c": 3})
        self.assertEqual(B.value_of("-"), (None, True))
        self.assertEqual(B.value_of(12), (12, False))
        self.assertEqual(B.slug("Habeas Corpus--General"), "habeas-corpus-general")
        self.assertEqual(B.slug(None), "none")

    def test_bankruptcy_court_ids_and_exceptions(self):
        bk = {"meb": 1, "arb": 1, "nebraskab": 1, "nysb": 1}
        self.assertEqual(B.bankruptcy_id("med", bk), "meb")
        self.assertEqual(B.bankruptcy_id("azd", bk), "arb")
        self.assertEqual(B.bankruptcy_id("ned", bk), "nebraskab")
        self.assertIsNone(B.bankruptcy_id("ard", bk, {"med": 1}))          # the combined Arkansas row is not a district: it must not become Arizona's arb
        self.assertIsNone(B.bankruptcy_id("ilnd", bk))                     # id not in the spine: not linked

    def test_series_keep_dashes_flags_and_missing_years_separate(self):
        jb = self.jb()
        values, dashes, flagged = jb.series("S-19", "med", "out_12m")
        self.assertEqual(values[-2:], [None, None])
        self.assertEqual(dashes, [2024, 2025])                              # printed dash: no value, listed
        self.assertEqual(flagged, [2024, 2025])                             # the row was part of a group that did not add up
        self.assertEqual(values[0], None)                                   # a year without a file is Not recorded (None), not zero
        self.assertEqual(jb.failed_checks("S-19"), {2024: [{"check": "arith", "detail": "x"}]})

    def test_district_records_derive_rates_and_shares_from_printed_values(self):
        rows = B.district_records(self.ctx())
        by = {r["id"]: r["item"] for r in rows}
        med = by["med"]["series"]
        self.assertEqual(med["terminated_total"][-2:], [1000, 1100])
        self.assertEqual(med["trial_rate_pct"][-1], round(10 / 1100 * 100, 2))
        self.assertEqual(med["share_of_national_terminations_pct"][-1], round(1100 / 1500 * 100, 2))
        self.assertEqual(med["mdl_pending_share_of_national_pct"][-1], round(30 / 50 * 100, 2))
        self.assertEqual(by["med"]["ranks"]["terminated_total"], 1)
        self.assertEqual(by["nhd"]["ranks"]["terminated_total"], 2)
        self.assertEqual(by["med"]["dash"]["mdl_out_12m"], [2024, 2025])
        self.assertEqual(by["med"]["unreconciled_years"]["mdl_out_12m"], [2024, 2025])
        self.assertIsNone(med["terminated_total"][0])                      # FY2005 has no file in the fixture
        badges = {r["id"]: r["item"]["badges"] for r in rows}
        self.assertEqual(badges["med"], ["Source total mismatch"])
        self.assertEqual(badges["nhd"], [])

    def test_bankruptcy_records_show_dashes_and_not_recorded_and_keep_the_combined_arkansas_row_apart(self):
        f2 = lambda total, ch7, ch11, ch13: {f: 0 for f in B.F2_FIELDS} | {"total": total, "ch7": ch7, "ch11": ch11, "ch13": ch13, "bus_total": 1}  # noqa: E731
        jb = B.JB.__new__(B.JB)
        jb.tables = {("C-4A", 2025): payload("C-4A", 2025, [trow({})])}
        jb.tables[("F-2", 2024)] = payload("F-2", 2024, [trow({"total": 300}), drow("med", "ca1", f2(200, 150, 5, 45)), drow("ared", "ca8", f2(40, 20, 0, 20)), drow("arwd", "ca8", f2(60, 30, 0, 30))])
        jb.tables[("F-2", 2025)] = payload("F-2", 2025, [trow({"total": 280}), drow("med", "ca1", f2(170, 120, D, 50)),
                                                        {"label": "AR", "kind": "district", "court_id": "ard", "circuit": "ca8", "fields": f2(60, 40, D, 20), "unreconciled": []}])
        spine = {"districts": {"med": ["District Court, D. Maine", "ME"], "ared": ["x", "AR"], "arwd": ["y", "AR"]},
                 "bankruptcy": {"meb": ["Bankruptcy Court, D. Maine", "ME"], "areb": ["Bankruptcy Court, E.D. Arkansas", "AR"], "arwb": ["Bankruptcy Court, W.D. Arkansas", "AR"]}}
        rows, unmapped = B.bankruptcy_records(B.Context(jb, spine, {"circuits": {}}, {"ids": ["ME", "AR"]}, []))
        by = {r["id"]: r["item"] for r in rows}
        self.assertEqual(unmapped, ["AR"])                                     # the combined row is not mapped to either Arkansas court (nor to Arizona's arb)
        self.assertEqual(sorted(by), ["areb", "arkansas-combined", "arwb", "meb"])
        self.assertEqual(by["areb"]["series"]["total"][YEAR_2024], 40)
        self.assertIsNone(by["areb"]["series"]["total"][-1])                  # no separate FY2025 row is printed: Not recorded, never zero
        self.assertEqual(by["areb"]["cells"]["filed"], "Not recorded")
        self.assertEqual(by["meb"]["cells"]["filed"], "170")
        self.assertEqual(by["meb"]["cells"]["ch11"], "–")                  # printed dash: shown as a dash, stored as no value, listed
        self.assertIsNone(by["meb"]["series"]["ch11"][-1])
        self.assertEqual(by["meb"]["dash"]["ch11"], [2025])
        combined = by["arkansas-combined"]
        self.assertEqual(combined["cells"]["filed"], "60")                     # its own figure, not the display copy of the court built before it
        self.assertEqual(combined["cells"]["ch7"], "40")
        self.assertEqual(combined["cells"]["ch11"], "–")
        self.assertEqual(combined["dash"]["ch11"], [2025])

    def test_nature_of_suit_keeps_only_labels_present_in_every_year_of_the_scheme(self):
        jb = B.JB.__new__(B.JB)
        jb.tables = {}
        for y in B.NOS_YEARS:
            f = {k: 5 for k in B.C4A_FIELDS}
            rows = [trow(dict(f, total=1000)), {"kind": "nos", "jurisdiction": "us", "category": "Contract Actions", "label": "Franchise", "occurrence": 1, "fields": dict(f, total=D, trial_total=D)}]
            if y > 2011:
                rows.append({"kind": "nos", "jurisdiction": "us", "category": "Tort Actions", "label": "Added Later", "occurrence": 1, "fields": f})
            rows.append({"kind": "nos", "jurisdiction": "local", "category": None, "label": "Contract Actions", "occurrence": 1, "fields": f})
            for occ in (1, 2):   # the same label printed twice in one group (memo lines): ambiguous, never published
                rows.append({"kind": "nos", "jurisdiction": "us", "category": "Actions Under Statutes", "label": "Consumer Credit", "occurrence": occ, "fields": f})
            jb.tables[("C-4", y)] = payload("C-4", y, rows)
        jb.tables[("C-4A", 2025)] = payload("C-4A", 2025, [trow({})])
        ctx = B.Context(jb, {"districts": {}, "bankruptcy": {}}, {"circuits": {}}, {"ids": []}, [])
        records, info = B.nos_records(ctx)
        ids = [r["id"] for r in records]
        self.assertEqual(info, {"stable_keys": 2, "all_keys": 5, "excluded_repeated_labels": 1})
        self.assertNotIn("nos-us-tort-actions-added-later", ids)
        self.assertFalse([i for i in ids if "consumer-credit" in i])
        franchise = next(r for r in records if r["id"].endswith("franchise"))
        self.assertEqual(franchise["item"]["dash"]["total"], B.NOS_YEARS)    # printed dashes: no value, never zero
        self.assertIsNone(franchise["item"]["series"]["total"][0])
        self.assertIsNone(franchise["item"]["series"]["trial_rate_pct"][0])
        local = next(r for r in records if "-local-" in r["id"])
        self.assertEqual(local["item"]["cells"]["group"], "Not grouped")


if __name__ == "__main__":
    unittest.main()
