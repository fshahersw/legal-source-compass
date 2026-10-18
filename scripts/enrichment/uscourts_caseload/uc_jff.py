"""Judicial Facts and Figures (JFF) year-row tables from the published .xlsx files (stream uscourts_caseload).

A JFF table is a sheet whose data rows start with a fiscal year ("2025", "1990 ²") and whose columns are measures under one to three
header rows (merged group headings). Each edition (2014, 2016-2025) lists the latest ten or so years plus anchor years (1990, 1995, ...);
the 2014 edition lists more. This module reads one sheet and then merges the editions of one table.

Rules kept everywhere:
  * every figure keeps the address of the cell it was read from (edition file, row, column), so a reader can open the workbook and compare;
  * values are copied as published; a printed dash is kept as a dash marker (stored as None with the year listed in `dash`), an empty cell is
    None ("Not recorded"); nothing is imputed;
  * a value is taken from the latest edition that lists the year; when an earlier edition printed a different number it is kept as a revision;
  * columns are matched between editions by their header text (letters only); a column an edition does not have is not filled from it.
"""
import re
import warnings
from pathlib import Path

import uc_common as U

YEAR = re.compile(r"^\s*((?:19|20)\d\d)(?!\d)")
MAX_COLS = 40
MAX_ROWS = 400
DASH = U.DASH


def col_letter(index):
    """0-based column index -> spreadsheet letters (0 -> A)."""
    out = ""
    index += 1
    while index:
        index, rem = divmod(index - 1, 26)
        out = chr(65 + rem) + out
    return out


def clean_label(text):
    """Header text without footnote markers (superscripts and a trailing standalone digit), single-spaced."""
    text = U.norm_text(text)
    text = "".join(ch for ch in text if ch not in U.SUPERSCRIPTS)
    text = re.sub(r"\s+\d$", "", text).strip()
    return re.sub(r"\s+", " ", text)


def signature(label):
    return re.sub(r"[^a-z]", "", label.lower())


def keyed(columns):
    """[(key, column)] where key is the header-text signature plus its occurrence number, so repeated headings (several 'Total' columns)
    stay distinct and are matched between editions by order of appearance."""
    seen = {}
    out = []
    for c in columns:
        sig = signature(c["label"])
        seen[sig] = seen.get(sig, 0) + 1
        out.append((f"{sig}#{seen[sig]}", c))
    return out


def unique_labels(columns):
    """Display labels, with the column letter added to labels that occur more than once."""
    counts = {}
    for c in columns:
        counts[c["label"]] = counts.get(c["label"], 0) + 1
    return [c["label"] if counts[c["label"]] == 1 else f"{c['label']} (column {c['letter']})" for c in columns]


def read_sheet(path):
    """Parse the first sheet: {"columns": [{"index", "letter", "label", "header": [...]}], "rows": {year: {"row": n, "mark": str|None, "cells": {index: value}}}, "notes": [...]}.
    Returns None when the sheet has no year rows."""
    import openpyxl

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    grid = [list(row) for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, MAX_ROWS), max_col=min(ws.max_column, MAX_COLS), values_only=True)]
    first = None
    for r, row in enumerate(grid):
        if row and row[0] is not None and YEAR.match(U.norm_text(row[0])) and any(cell_value(c)[0] is not None or cell_value(c)[1] for c in row[1:]):
            first = r
            break
    if first is None:
        return None
    # header rows: rows above the first data row with a distinct non-empty text beyond column 0 after the merges are filled
    filled = [list(row) for row in grid]
    for rng in ws.merged_cells.ranges:
        if rng.min_row - 1 >= first:
            continue
        value = grid[rng.min_row - 1][rng.min_col - 1] if rng.min_col - 1 < len(grid[rng.min_row - 1]) else None
        for r in range(rng.min_row - 1, min(rng.max_row, first)):
            for c in range(rng.min_col - 1, min(rng.max_col, MAX_COLS)):
                if r < len(filled) and c < len(filled[r]):
                    filled[r][c] = value
    header_rows = []
    for r in range(first):
        distinct = {U.norm_text(v) for v in filled[r][1:] if v is not None and U.norm_text(v)}
        if len(distinct) >= 2 or (len(distinct) == 1 and any(v is not None for v in grid[r][1:])):
            header_rows.append(r)
    data = {}
    notes = []
    ncols = max(len(row) for row in grid)
    for r in range(first, len(grid)):
        row = grid[r] + [None] * (ncols - len(grid[r]))
        text = U.norm_text(row[0]) if row[0] is not None else ""
        m = YEAR.match(text)
        if m and any(c is not None for c in row[1:]):
            mark = U.split_footnote(text)[1]
            data[int(m.group(1))] = {"row": r + 1, "mark": mark, "cells": {c: row[c] for c in range(1, ncols) if row[c] is not None}}
        elif text:
            notes.append(text)
    used = sorted({c for d in data.values() for c in d["cells"]})
    columns = []
    for c in used:
        header = []
        for r in header_rows:
            cell = filled[r][c] if c < len(filled[r]) else None
            label = clean_label(cell) if cell is not None else ""
            if label and (not header or header[-1] != label):
                header.append(label)
        columns.append({"index": c, "letter": col_letter(c), "label": " / ".join(header) or f"Column {col_letter(c)}", "header": header})
    return {"columns": columns, "rows": data, "notes": notes, "sheet": ws.title, "sheets": len(wb.worksheets)}


MARKED = re.compile(r"^\s*(-?\d[\d,]*(?:\.\d+)?)\s*(\*+|\[\d+\]|\s\d)\s*$")


def cell_value(raw):
    """Published cell -> (value, dash, mark). A number; (None, True, None) for a printed dash; (None, False, None) for an unreadable cell.
    A number printed with a footnote marker ('302*', '730 [4]', '1,399,496 2', '663 6') is read as the number and the marker is returned
    (a trailing single digit after a space counts as a marker only when the number has three or more digits)."""
    v = U.to_value(raw)
    if v == DASH:
        return None, True, None
    if v is not None:
        return v, False, None
    if isinstance(raw, str):
        m = MARKED.match(U.norm_text(raw))
        if m:
            number, mark = m.group(1), m.group(2).strip()
            if not (mark.isdigit() and len(number.replace(",", "").split(".")[0]) < 3):
                return U.to_value(number), False, mark
    return None, False, None


def merge_editions(parsed):
    """parsed: [(edition, source dict, sheet dict)] oldest first. Returns the merged table:
    {"columns": [{label, letter}], "years": [...], "values": {label: [per year]}, "dash": {label: [years]}, "cells": {year: {label: {edition, address}}},
     "year_edition": {year: edition}, "revisions": [...], "skipped": [...], "notes": [...]}"""
    latest = parsed[-1][2]
    canon = latest["columns"]
    names = dict(zip((k for k, _ in keyed(canon)), unique_labels(canon)))
    canon_keys = list(names)
    years = sorted({y for _, _, s in parsed for y in s["rows"]})
    values = {names[k]: {} for k in canon_keys}
    dash = {names[k]: [] for k in canon_keys}
    cells = {}
    year_edition = {}
    revisions = []
    skipped = []
    for edition, source, sheet in parsed:  # oldest first: later editions overwrite, differences are kept
        by_key = dict(keyed(sheet["columns"]))
        matched = {names[k]: by_key[k] for k in canon_keys if k in by_key}
        if len(matched) < len(canon):
            skipped.append({"edition": edition, "missing_columns": [names[k] for k in canon_keys if k not in by_key]})
        for year, entry in sheet["rows"].items():
            for label, col in matched.items():
                raw = entry["cells"].get(col["index"])
                if raw is None:
                    continue
                v, is_dash, mark = cell_value(raw)
                prior = values[label].get(year)
                if prior is not None and v is not None and prior["value"] != v:
                    revisions.append({"year": year, "column": label, "earlier": prior["value"], "earlier_edition": prior["edition"], "later": v, "later_edition": edition})
                values[label][year] = {"value": v, "dash": is_dash, "mark": mark, "edition": edition, "address": f"{col['letter']}{entry['row']}", "file": source["path"]}
            year_edition[year] = edition
    out_values, out_dash, out_marks = {}, {}, {}
    for label in values:
        marks = {y: values[label][y]["mark"] for y in years if values[label].get(y, {}).get("mark")}
        if marks:
            out_marks[label] = marks
        out_values[label] = [values[label].get(y, {}).get("value") for y in years]
        d = [y for y in years if values[label].get(y, {}).get("dash")]
        if d:
            out_dash[label] = d
    for y in years:
        cells[y] = {label: {"edition": values[label][y]["edition"], "address": values[label][y]["address"]} for label in values if y in values[label]}
    return {"columns": [{"label": n, "letter": c["letter"]} for n, c in zip(unique_labels(canon), canon)], "years": years, "values": out_values, "dash": out_dash, "marks": out_marks, "cells": cells,
            "year_edition": year_edition, "revisions": revisions, "skipped": skipped, "notes": latest["notes"]}
