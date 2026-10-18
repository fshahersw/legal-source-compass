"""Independent spot-check of published court_stats records against freshly downloaded official tables.

Each check re-downloads the official U.S. Courts file (not the cached raw copy), reads one cell with a plain
openpyxl / pdfplumber scan written independently of tables.py/extract.py, and compares it with the value
stored in Supabase. Prints one line per check: dataset/record field: published=X source=Y url.
Usage: python3 verify_published.py
"""
import io
import json
import os
import re
import sys
import urllib.request

import openpyxl

UA = "SeegerWeiss Research firasshaher99@gmail.com"
# (dataset, record id, item path, fiscal-year index or None, file, sheet substring, row label, column header regex, after-label, digits)
CHECKS = [
    ("court_stats_district_nos", "njd", "stats.pl_total", 2025, ("C-11", 2025), "C-11", "NJ", r"^2025 Total$", None, None),
    ("court_stats_district_nos", "flnd", "stats.pl_total", 2020, ("C-11", 2020), "C-11", "FL,N", r"^2020 Total$", None, None),
    ("court_stats_district_nos", "njd", "stats.civil_filed", 2025, ("C", 2025), "", "NJ", r"Filed / 2025", None, None),
    ("court_stats_district_nos", "ilnd", "stats.judgeships", 2025, ("X-1A", 2025), "", "IL,N", r"(^|/ )Judgeships$", None, None),
    ("court_stats_district_times", "njd", "times.median_total", 2025, ("C-5", 2025), "", "NJ", r"Total Cases / Median", None, 1),
    ("court_stats_district_times", "laed", "times.pending_over_3y", 2025, ("C-6", 2025), "", "LA,E", r"3 Years or More / Number", None, None),
    ("court_stats_circuits", "ca9", "circuit.appeals_filed", 2025, ("B", 2025), "", "9th", r"Filed / 2025", None, None),
    ("court_stats_circuits", "ca3", "circuit.median_notice_to_final", 2025, ("B-4", 2025), "", "3rd", r"Notice of Appeal or Docket Date to Last Opinion.*Interval", None, 1),
    ("court_stats_circuits", "ca11", "circuit.diversity_pi_pl_appeals", 2025, ("B-7", 2025), "", "Personal Injury/Product Liability", r"11th", "Diversity of Citizenship", None),
    ("court_stats_national", "summary:national", "summary.series.civil_filed", 2025, ("C", 2025), "", "Total", r"Filed / 2025", None, None),
    ("court_stats_national", "summary:national", "summary.series.pl_total", 2020, ("C-11", 2020), "C-11", "Total", r"^2020 Total$", None, None),
    ("court_stats_national", "fy2023", "national.c5.median_total", None, ("C-5", 2023), "", "Total", r"Total Cases / Median", None, 1),
]
PDF_CHECKS = [
    # FY2016 is published as PDF only: the District of Columbia row of Table C-11 (2015 total, 2016 total, ...).
    ("court_stats_district_nos", "dcd", "stats.pl_total", 2016, ("C-11", 2016), "DC", 1),
]


SOURCES = {}


def source_url(ref):
    """Official file URL for (table, fiscal year) as stored in the published national summary."""
    if not SOURCES:
        SOURCES.update(published("court_stats_national", "summary:national", "summary.sources")[0])
    table, year = ref
    return SOURCES[table][str(year)]["url"]


def download(ref):
    request = urllib.request.Request(source_url(ref), headers={"User-Agent": UA})
    with urllib.request.urlopen(request, timeout=120) as response:
        return response.read()


def norm(v):
    return re.sub(r"\s+", " ", str(v or "").replace("_x000D_", " ")).strip()


def xlsx_cell(data, sheet_hint, label, header, after):
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    sheets = [ws for ws in wb.worksheets if sheet_hint in ws.title] or wb.worksheets
    for ws in sheets:
        grid = [list(r) for r in ws.iter_rows(values_only=True)]
        for rng in ws.merged_cells.ranges:
            v = ws.cell(rng.min_row, rng.min_col).value
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    grid[r - 1][c - 1] = v
        start = 0
        if after:
            start = next((i for i, r in enumerate(grid) if r and norm(r[0]).startswith(after)), None)
            if start is None:
                continue
        row_i = next((i for i in range(start, len(grid)) if grid[i] and re.sub(r"[\d¹²³]+$", "", norm(grid[i][0])) == label), None)
        if row_i is None:
            continue
        # heading rows: everything above the first data row ("Total")
        top = next(i for i, r in enumerate(grid) if r and norm(r[0]).upper() == "TOTAL")
        width = max(len(r) for r in grid)
        for col in range(1, width):
            text = " / ".join(dict.fromkeys(norm(grid[r][col]) for r in range(1, top) if col < len(grid[r]) and norm(grid[r][col])))
            if re.search(header, text):
                return grid[row_i][col], text
    return None, None


def pdf_row(data, label):
    import pdfplumber

    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            for line in (page.extract_text() or "").splitlines():
                if re.match(rf"^{re.escape(label)}\s+[\d,]", line.strip()):
                    return line.strip()
    return None


def published(dataset, rid, path):
    base = os.environ["EXTERNAL_SUPABASE_URL"].rstrip("/") + "/rest/v1/"
    key = os.environ["EXTERNAL_SUPABASE_KEY"]
    head, *rest = path.split(".")
    request = urllib.request.Request(base + f"corpus_records?select=item&dataset=eq.{dataset}&id=eq.{rid}",
                                     headers={"apikey": key, "Authorization": f"Bearer {key}"})
    rows = json.loads(urllib.request.urlopen(request, timeout=60).read())
    value = rows[0]["item"][head]
    for part in rest:
        value = value[part]
    return value, rows[0]["item"]


def main():
    cache = {}
    ok = True
    for dataset, rid, path, year, name, sheet, label, header, after, digits in CHECKS:
        data = cache.get(name) or cache.setdefault(name, download(name))
        source, heading = xlsx_cell(data, sheet, label, header, after)
        value, item = published(dataset, rid, path)
        if isinstance(value, list):
            years = item.get("years") or item.get(path.split(".")[0], {}).get("years") or (item.get("summary") or {}).get("years")
            value = value[years.index(year)]
        src = round(float(source), digits) if digits is not None and source is not None else source
        match = src == value
        ok &= match
        print(f"{'OK ' if match else 'BAD'} {dataset}/{rid} {path}{f' FY{year}' if year else ''}: published={value} source={src} ({heading}) {source_url(name)}")
    for dataset, rid, path, year, name, label, position in PDF_CHECKS:
        line = pdf_row(download(name), label)
        value, item = published(dataset, rid, path)
        value = value[item["stats"]["years"].index(year)]
        tokens = line.split()[1:] if line else []
        src = int(tokens[position].replace(",", "")) if len(tokens) > position else None
        match = src == value
        ok &= match
        print(f"{'OK ' if match else 'BAD'} {dataset}/{rid} {path} FY{year}: published={value} source={src} (PDF line {line!r}) {source_url(name)}")
    sys.exit(0 if ok else 1)


if __name__ == "__main__":
    main()
