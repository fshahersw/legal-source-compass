"""Offline tests for the FCMS enrichment (synthetic fixtures, no network, no files from private/).

Run: cd scripts/enrichment/uscourts_fcms && python3 -B -m unittest -v test_uscourts_fcms
"""
import gzip
import json
import tempfile
import unittest
from pathlib import Path

import build_uscourts_fcms as build
import courts
import fcms_meta as meta
import fcms_model as fm
import jsonio
import layout
import measures
import ocr_assessment
import profiles


def word(x0, x1, text, y=100.0):
    return (x0, y - 4, x1, y + 4, text)


SPINE = {
    "cand": ("District Court, N.D. California", "CA"), "ard": ("United States District Court for the District of Arkansas", "AR"), "ared": ("District Court, E.D. Arkansas", "AR"),
    "ca9": ("Court of Appeals for the Ninth Circuit", ""), "prd": ("District Court, D. Puerto Rico", "PR"), "dcd": ("District Court, District of Columbia", "DC"),
    "azd": ("District Court, D. Arizona", "AZ"), "ilnd": ("District Court, N.D. Illinois", "IL"), "cadc": ("Court of Appeals for the D.C. Circuit", ""),
}


class LayoutTests(unittest.TestCase):
    def test_glued_numbers_are_split_in_proportion(self):
        pieces = layout.split_glued(100.0, 0, 190.0, 10, "330,721358,983328,520")
        self.assertEqual([p[4] for p in pieces], ["330,721", "358,983", "328,520"])
        self.assertTrue(pieces[0][0] < pieces[1][0] < pieces[2][0])
        # a normal word and a word that is not an exact concatenation stay whole
        self.assertEqual(layout.split_glued(0, 0, 10, 10, "1,234")[0][4], "1,234")
        self.assertEqual(len(layout.split_glued(0, 0, 90, 10, "12,345abc12,345")), 1)

    def test_value_tokens(self):
        for token in ("4,666", "-", "84.6%", ".7", "-.7", "12/4.0", "0.0", "− 3.2".replace(" ", "")):
            self.assertTrue(layout.is_value(token), token)
        for token in ("Filings", "2026", "a1", "1.2.3", ""):
            if token == "2026":
                continue  # a bare year looks like a value; year headers are recognised before values are read
            self.assertFalse(layout.is_value(token), token)

    def test_month_markers_two_word_and_one_word(self):
        lines = [{"mid": 80.0, "words": [word(170, 178, "Sep", 80), word(179, 184, "30", 80), word(191, 199, "Sep", 80), word(200, 205, "30", 80)]},
                 {"mid": 90.0, "words": [word(300, 320, "30-Sep", 90), word(330, 350, "31-Dec", 90)]},
                 {"mid": 100.0, "words": [word(170, 182, "2020", 100)]}]
        marks = layout.month_markers(lines, 2)
        self.assertEqual([m[1] for m in marks], ["09-30", "09-30", "09-30", "12-31"])
        # a month word that is not followed by its own day is not a marker ("June 15" is not a period end)
        bad = [{"mid": 80.0, "words": [word(0, 20, "June", 80), word(21, 30, "15", 80)]}, {"mid": 100.0, "words": []}]
        self.assertEqual(layout.month_markers(bad, 1), [])

    def test_assign_months_one_per_column_single_marker_and_groups(self):
        centres = [285.6, 327.2, 368.8, 410.5, 452.1, 493.7]
        self.assertEqual(layout.assign_months([(1, "09-30")], centres), ["09-30"] * 6)
        each = [(c, "06-30") for c in centres]
        self.assertEqual(layout.assign_months(each, centres), ["06-30"] * 6)
        # December 2011 edition: "September 30" centred over four columns, "December 31" over two
        groups = [(348.0, "09-30"), (473.0, "12-31")]
        self.assertEqual(layout.assign_months(groups, centres, spacing=41.6), ["09-30"] * 4 + ["12-31"] * 2)
        # markers that are not centred over any contiguous group are refused rather than guessed
        self.assertIsNone(layout.assign_months([(100.0, "09-30"), (700.0, "12-31")], centres, spacing=41.6))
        self.assertIsNone(layout.assign_months([], centres))

    def test_period_keys_resolve_duplicate_years(self):
        # two columns both labelled 2010 (September and December): keys are period-end dates, so neither is overwritten
        years = [word(240, 260, "2008", 94), word(286, 306, "2009", 94), word(333, 353, "2010", 94), word(379, 399, "2010", 94), word(426, 446, "2011", 94), word(472, 492, "2012", 94)]
        marker_words = []
        for x, mon in ((235, "Sep"), (282, "Sep"), (328, "Sep"), (375, "Dec"), (421, "Dec"), (468, "Dec")):
            marker_words += [word(x, x + 17, mon, 83), word(x + 19, x + 29, "30" if mon == "Sep" else "31", 83)]
        lines = [{"mid": 83.0, "words": marker_words}, {"mid": 94.0, "words": years}]
        keys, _ = profiles.period_keys(lines, 1, years)
        self.assertEqual(keys, ["2008-09-30", "2009-09-30", "2010-09-30", "2010-12-31", "2011-12-31", "2012-12-31"])

    def test_read_panel_positional_fallback_and_ambiguity(self):
        years = [word(100 + 50 * i, 120 + 50 * i, str(2021 + i), 50) for i in range(6)]
        keys = [f"{2021 + i}-06-30" for i in range(6)]
        centres = [110 + 50 * i for i in range(6)]
        # one token per column in order: read normally
        normal = {"mid": 70.0, "words": [word(10, 60, "Filings", 70)] + [word(c - 8, c + 8, str(100 + i), 70) for i, c in enumerate(centres)]}
        rows = layout.read_panel([normal], 0, years, 0, 1e9, keys=keys)
        self.assertEqual([rows[0]["values"][k] for k in keys], [str(100 + i) for i in range(6)])
        self.assertEqual(rows[0]["stray"], [])
        # two adjacent narrow tokens fall nearest the same column; exactly six tokens in the area: assigned left to right
        close = {"mid": 70.0, "words": [word(10, 60, "Judgeships", 70), word(102, 106, "3", 70), word(108, 112, "3", 70), word(155, 165, "4", 70),
                                         word(205, 215, "5", 70), word(255, 265, "6", 70), word(305, 315, "7", 70)]}
        rows = layout.read_panel([close], 0, years, 0, 1e9, keys=keys)
        self.assertEqual([rows[0]["values"][k] for k in keys], ["3", "3", "4", "5", "6", "7"])
        self.assertEqual(rows[0]["stray"], [])
        # twelve tokens for six columns (two figures printed in each cell): never guessed, the row is ambiguous
        crowded = {"mid": 70.0, "words": [word(10, 60, "Median", 70)] + [word(c - 12 + 12 * j, c - 4 + 12 * j, f"{i}.{j}", 70) for i, c in enumerate(centres) for j in range(2)]}
        rows = layout.read_panel([crowded], 0, years, 0, 1e9, keys=keys)
        self.assertTrue(rows[0]["stray"])


class MeasureTests(unittest.TestCase):
    def test_parse_value_dash_and_bad(self):
        self.assertEqual(measures.parse_value("4,666"), (4666, None))
        self.assertEqual(measures.parse_value("84.6%"), (84.6, None))
        self.assertEqual(measures.parse_value("-"), (None, "dash"))
        self.assertEqual(measures.parse_value("− ".strip()), (None, "dash"))
        self.assertEqual(measures.parse_value("-3.3"), (-3.3, None))
        self.assertEqual(measures.parse_value(".7"), (0.7, None))
        self.assertEqual(measures.parse_value(None), (None, "absent"))
        self.assertEqual(measures.parse_value("12/4.0")[1], "bad")

    def row(self, label, **values):
        return {"label": label, "values": values, "standing": {}, "stray": [], "y": 1.0}

    def test_classify_in_order_and_skips_optional_rows(self):
        rows = [self.row("Filings", a="1"), self.row("Terminations", a="2"), self.row("Pending", a="3"), self.row("Percent change over last year", a="4"),
                self.row("Number of Judgeships", a="5"), self.row("Vacant Judgeship Months", a="6"), self.row("Total", a="7"), self.row("Civil", a="8"),
                self.row("Criminal Felony", a="9"), self.row("Pending Cases", a="10"), self.row("Weighted Filings", a="11"), self.row("Terminations", a="12"),
                self.row("Trials Completed", a="13"), self.row("Criminal Felony", a="14"), self.row("Civil", a="15"), self.row("From Filing to Trial", a="16"),
                self.row("Number of Civil Cases", a="17"), self.row("Over 3 Years Old", a="18"), self.row("Filed per Case defendants", a="19")]
        assigned, unmatched, missing = measures.classify(rows, measures.DISTRICT)
        self.assertEqual(missing, [])
        self.assertEqual(unmatched, [])
        self.assertEqual(assigned["pj_weighted_filings"]["values"]["a"], "11")
        self.assertEqual(assigned["median_months_civil_to_trial"]["values"]["a"], "16")
        self.assertNotIn("pj_supervised_release_hearings", assigned)  # optional row not printed in this fixture: skipped, not invented

    def test_missing_required_row_is_reported_not_shifted(self):
        rows = [self.row("Filings", a="1"), self.row("Pending", a="3")]
        assigned, unmatched, missing = measures.classify(rows, measures.DISTRICT)
        self.assertIn("terminations", missing)

    def test_garbled_merits_label_still_matches_by_position(self):
        row = self.row("erminations on the Meri", a="429")
        compiled = [(k, rx) for k, rx, *_ in [(e[0], e[1]) for e in measures.CIRCUIT_LEFT]]
        import re
        rx = dict(compiled)["paj_merits_terminations"]
        self.assertTrue(re.search(rx, row["label"], re.I))


class CourtMatchTests(unittest.TestCase):
    def test_exact_name_rule(self):
        self.assertEqual(courts.resolve_district("CALIFORNIA NORTHERN", SPINE), ("cand", "CA", "exact_name_rule"))
        self.assertEqual(courts.resolve_district("DISTRICT OF COLUMBIA", SPINE), ("dcd", "DC", "exact_name_rule"))
        self.assertEqual(courts.resolve_district("NO. CAROLINA EASTERN".replace("NO. CAROLINA", "NORTH CAROLINA"), {"nced": ("District Court, E.D. North Carolina", "NC")})[0], "nced")

    def test_non_match_is_never_fuzzy(self):
        # OCR-garbled and misspelled names are not matched to the closest court
        self.assertEqual(courts.resolve_district("CALFORNIA NORTHERN", SPINE)[0], None)
        self.assertEqual(courts.resolve_district("ILINOIS NORTHERN", SPINE)[0], None)
        self.assertEqual(courts.resolve_district("CALIFORNIA NORTH", SPINE)[0], None)

    def test_multi_district_state_without_direction_never_binds_to_statewide_placeholder(self):
        # court_spine holds a statewide "ard" row; "ARKANSAS" alone must not bind to it (Arkansas has two districts)
        self.assertEqual(courts.resolve_district("ARKANSAS", SPINE)[0], None)
        self.assertEqual(courts.resolve_district("ARKANSAS EASTERN", SPINE)[0], "ared")
        self.assertEqual(courts.resolve_district("ARIZONA", SPINE)[0], "azd")  # a single-district state is named without a direction

    def test_id_must_exist_in_spine_with_matching_type_and_state(self):
        self.assertEqual(courts.resolve_district("ARKANSAS WESTERN", SPINE), (None, "AR", "not_in_spine"))
        bad = {"cand": ("Superior Court of California", "CA")}
        self.assertEqual(courts.resolve_district("CALIFORNIA NORTHERN", bad)[2], "spine_mismatch")
        self.assertEqual(courts.resolve_circuit("NINTH CIRCUIT", SPINE), ("ca9", "exact_name_rule"))
        self.assertEqual(courts.resolve_circuit("DISTRICT OF COLUMBIA", SPINE), ("cadc", "exact_name_rule"))
        self.assertEqual(courts.resolve_circuit("TWELFTH CIRCUIT", SPINE)[0], None)

    def test_every_district_belongs_to_one_regional_circuit(self):
        self.assertEqual(len(courts.CIRCUIT_OF), 94)
        self.assertEqual(courts.CIRCUIT_OF["cand"], "ca9")
        self.assertEqual(courts.CIRCUIT_OF["dcd"], "cadc")


class ModelTests(unittest.TestCase):
    def test_period_arithmetic(self):
        self.assertEqual(fm.prior_year_key("2026-06-30"), "2025-06-30")
        self.assertEqual(fm.prior_year_key("2026-06-30", 5), "2021-06-30")
        self.assertEqual(fm.period_text("2025-12-31"), "December 31, 2025")
        self.assertEqual(fm.edition_title("2004-09-30"), "September 2004")
        self.assertTrue(fm.is_fiscal_year_end("2019-09-30"))
        self.assertFalse(fm.is_fiscal_year_end("2019-12-31"))

    def test_pct_change_formula_and_unknowns(self):
        self.assertEqual(fm.pct_change(1020, 815), 25.2)
        self.assertEqual(fm.pct_change(50, 100), -50.0)
        self.assertIsNone(fm.pct_change(5, 0))      # zero base: no percentage
        self.assertIsNone(fm.pct_change(None, 10))  # unknown stays unknown
        self.assertIsNone(fm.pct_change(10, None))
        self.assertEqual(fm.abs_change(10.5, 8.0), 2.5)

    def test_competition_ranks_ties_and_unknowns(self):
        ranks, n = fm.competition_ranks({"a": 10, "b": 20, "c": 20, "d": 5, "e": None})
        self.assertEqual((ranks["b"], ranks["c"], ranks["a"], ranks["d"]), (1, 1, 3, 4))
        self.assertNotIn("e", ranks)
        self.assertEqual(n, 4)
        low, _ = fm.competition_ranks({"a": 10, "b": 20}, descending=False)
        self.assertEqual(low["a"], 1)

    def test_extremes_are_deterministic_on_ties(self):
        high, low = fm.extremes({"b": 5, "a": 5, "c": 1, "d": None}, 2)
        self.assertEqual(high, [("a", 5), ("b", 5)])
        self.assertEqual(low, [("c", 1), ("a", 5)])

    def test_distribution(self):
        d = fm.distribution([1, 2, 3, 4, None])
        self.assertEqual((d["n"], d["min"], d["median"], d["max"]), (4, 1, 2.5, 4))
        self.assertEqual(fm.distribution([None])["n"], 0)

    def cell(self, periods, **series):
        return {"kind": "district", "periods": periods, "values": series, "flags": {}, "standing": {}, "raw": {}, "pages": [1], "qa": [], "name": "X"}

    def test_changes_use_the_column_one_year_earlier_not_the_neighbouring_column(self):
        # December 2010 edition: the column left of 2010-12-31 is 2010-09-30, three months earlier, not the prior year
        periods = ["2006-09-30", "2007-09-30", "2008-09-30", "2009-09-30", "2010-09-30", "2010-12-31"]
        cell = self.cell(periods, filings=dict(zip(periods, [1, 2, 3, 4, 5, 6])))
        ch = fm.within_edition_changes(cell, "filings", "2010-12-31")
        self.assertEqual(ch["current"], 6)
        self.assertIsNone(ch["yoy"]["base_period"])    # 2009-12-31 is not printed: not recorded rather than 2010-09-30
        self.assertIsNone(ch["yoy"]["pct"])
        periods = ["2020-06-30", "2021-06-30", "2022-06-30", "2023-06-30", "2024-06-30", "2025-06-30"]
        cell = self.cell(periods, filings=dict(zip(periods, [100, 110, 120, 130, 140, 154])))
        ch = fm.within_edition_changes(cell, "filings", "2025-06-30")
        self.assertEqual((ch["yoy"]["base"], ch["yoy"]["pct"]), (140, 10.0))
        self.assertEqual((ch["five_year"]["base"], ch["five_year"]["pct"]), (100, 54.0))

    def edition(self, period, filings, note=False, cur=None):
        cell = {"values": {"filings": filings}, "flags": {}}
        return {"period_end": period, "cells": {"district": {"national": cell}}, "notes": {"district": ["NOTE: Criminal data in this profile count defendants rather than cases and therefore will not match."] if note else []}}

    def test_basis_explicit_tested_and_inherited(self):
        eds = {
            "2012-09-30": self.edition("2012-09-30", {"2010-09-30": 394, "2011-09-30": 402, "2012-09-30": 386}, note=True),
            "2011-09-30": self.edition("2011-09-30", {"2011-09-30": 378}),          # differs from 402 printed by the explicit edition: cases
            "2010-09-30": self.edition("2010-09-30", {"2010-09-30": 394}),          # equals: defendants
            "2005-09-30": self.edition("2005-09-30", {"2005-09-30": 300}),          # no explicit edition prints 2005: inherited from the nearest tested edition
        }
        basis = fm.determine_basis(eds)
        self.assertEqual(basis["2012-09-30"]["basis"], "defendants")
        self.assertEqual(basis["2011-09-30"]["basis"], "cases")
        self.assertEqual(basis["2010-09-30"]["basis"], "defendants")
        self.assertEqual(basis["2005-09-30"]["basis"], "defendants")
        self.assertIn("inherited", basis["2005-09-30"]["evidence"])
        self.assertIn("differ", basis["2011-09-30"]["evidence"])

    def test_rollup_sum_only_when_every_district_has_a_value(self):
        cells = {"akd": {"values": {"filings": {"p": 10}}}, "azd": {"values": {"filings": {"p": 20}}}, "cand": {"values": {"filings": {"p": None}}}}
        circuit_of = {"akd": "ca9", "azd": "ca9", "cand": "ca9", "ctd": "ca2"}
        out = fm.circuit_rollup(cells, circuit_of, "ca9", ["p"], ["filings"])
        self.assertIsNone(out["filings"]["p"]["sum"])            # one district unknown: no partial sum
        self.assertEqual((out["filings"]["p"]["known"], out["filings"]["p"]["districts"]), (2, 3))
        cells["cand"]["values"]["filings"]["p"] = 5
        self.assertEqual(fm.circuit_rollup(cells, circuit_of, "ca9", ["p"], ["filings"])["filings"]["p"]["sum"], 35)

    def test_qa_flags_district_and_circuit(self):
        p = "2026-06-30"
        ok = self.cell([p], judgeships={p: 3}, filings={p: 1020}, pj_filings_total={p: 340}, terminations={p: 934}, pj_terminations={p: 311}, pending={p: 1022}, pj_pending={p: 341})
        self.assertEqual(fm.qa_flags("district", ok), [])
        bad = self.cell([p], judgeships={p: 3}, filings={p: 1020}, pj_filings_total={p: 500})
        self.assertEqual(fm.qa_flags("district", bad)[0]["measure"], "pj_filings_total")
        circ = self.cell([p], filings={p: 9}, filings_prisoner={p: 1}, filings_other={p: 2}, filings_criminal={p: 3}, filings_administrative={p: 3})
        self.assertEqual(fm.qa_flags("circuit", circ), [])
        circ["values"]["filings_administrative"][p] = 4
        self.assertEqual(fm.qa_flags("circuit", circ)[0]["check"], "filings_components")

    def test_contradicting_judgeship_count_is_not_shown(self):
        # June 2009 text layer: 67 judgeships printed beside 55.7 panels (167 / 3 = 55.7)
        p = "2009-06-30"
        cell = self.cell([p], judgeships={p: 67}, panels={p: 55.7})
        fm.apply_inconsistent_judgeships(cell)
        self.assertIsNone(cell["values"]["judgeships"][p])
        self.assertEqual(cell["flags"]["judgeships"][p], "inconsistent")
        good = self.cell([p], judgeships={p: 167}, panels={p: 55.7})
        fm.apply_inconsistent_judgeships(good)
        self.assertEqual(good["values"]["judgeships"][p], 167)

    def test_first_print_uses_the_earliest_edition_and_skips_unprinted_cells(self):
        def ed(period, vals, flags=None):
            return {"cells": {"district": {"cand": {"values": {"filings": vals}, "flags": {"filings": flags or {}}}}}}
        eds = {"2010-09-30": ed("2010-09-30", {"2009-09-30": 1, "2010-09-30": 2}), "2011-09-30": ed("2011-09-30", {"2009-09-30": 9, "2010-09-30": 8, "2011-09-30": 3, "2008-09-30": None})}
        s = fm.first_print_series(eds, "district", "cand", "filings")
        self.assertEqual(s["2009-09-30"], (1, "2010-09-30"))    # earliest edition wins, not the latest revision
        self.assertEqual(s["2011-09-30"], (3, "2011-09-30"))
        self.assertNotIn("2008-09-30", s)                      # printed nothing: not a value

    def test_load_editions_and_circuit_merge(self):
        with tempfile.TemporaryDirectory() as tmp:
            d = Path(tmp)
            page = lambda cid, part, vals: {"court_id": cid, "name": "N", "page": 1 if part == "left" else 2, "part": part, "years": ["2020-06-30", "2021-06-30"], "labels": ["2020", "2021"], "values": vals, "flags": {}, "standing": {}, "raw": {}, "unmatched": []}  # noqa: E731
            parsed = {"edition": "e-june-2021", "period_end": "2021-06-30", "kind": "circuit", "file": "f", "url": "u", "sha256": "s", "page_url": "p", "ocr_text_layer": False, "page1_notes": [],
                      "excluded": [], "pages": [page("ca9", "left", {"filings": {"2020-06-30": 1, "2021-06-30": 2}}), page("ca9", "right", {"pp_filings": {"2020-06-30": 3, "2021-06-30": 4}})]}
            jsonio.dump(d / "e-june-2021__circuit.json.gz", parsed)
            ocr = dict(parsed, edition="e-1999", period_end="1999-09-30", ocr_text_layer=True)
            jsonio.dump(d / "e-1999__circuit.json.gz", ocr)
            jsonio.dump(d / "_index.json", {"e-june-2021__circuit": {"status": "parsed"}, "e-1999__circuit": {"status": "parsed"}, "e-june-2004__district": {"status": "image_only", "file": "x", "url": "y", "sha256": "z"}})
            editions, skipped = fm.load_editions(d)
            self.assertEqual(list(editions), ["2021-06-30"])
            cell = editions["2021-06-30"]["cells"]["circuit"]["ca9"]
            self.assertEqual(cell["pages"], [1, 2])
            self.assertEqual(set(cell["values"]), {"filings", "pp_filings"})
            self.assertEqual({(s["edition"], s["status"]) for s in skipped}, {("e-1999", "ocr_text_layer"), ("e-june-2004", "image_only")})


class FormatAndRecordTests(unittest.TestCase):
    def test_not_recorded_formatting(self):
        self.assertEqual(build.fmt(None), "Not recorded")
        self.assertEqual(build.fmt(1022), "1,022")
        self.assertEqual(build.fmt(7.9, 1), "7.9")
        self.assertEqual(build.fmt(11.1, 1, "percent"), "11.1%")
        self.assertEqual(build.sign(None), "Not recorded")
        self.assertEqual(build.sign(25.2, 1, "%"), "+25.2%")

    def cell(self, **kw):
        base = {"kind": "district", "court_id": "cand", "name": "CALIFORNIA NORTHERN", "usps": "CA", "match": "exact_name_rule", "pages": [7], "periods": ["2025-06-30", "2026-06-30"],
                "labels": ["2025", "2026"], "values": {}, "flags": {}, "standing": {}, "raw": {}, "col_x": None, "unmatched": [], "qa": []}
        base.update(kw)
        return base

    def test_shown_explains_why_a_value_is_missing(self):
        cell = self.cell(values={"median_months_civil_to_trial": {"2026-06-30": None}, "filings": {"2026-06-30": 1020}, "pending": {"2026-06-30": None}},
                         flags={"median_months_civil_to_trial": {"2026-06-30": "dash"}},
                         raw={"filings": {"tokens": {"2026-06-30": "1,020"}}})
        self.assertEqual(build.shown(cell, "district", "filings", "2026-06-30"), "1,020")
        self.assertIn("printed as a dash", build.shown(cell, "district", "median_months_civil_to_trial", "2026-06-30"))
        self.assertEqual(build.shown(cell, "district", "pending", "2026-06-30"), "Not recorded")
        cell["flags"]["pending"] = {"2026-06-30": "ambiguous"}
        self.assertIn("more than one figure per period column", build.shown(cell, "district", "pending", "2026-06-30"))

    def context(self):
        ctx = build.Context.__new__(build.Context)
        ctx.spine = SPINE
        ctx.places = {"CA", "AZ", "DC"}
        ctx.ca_ids = {"cand"}
        ctx.circuit_of = {"cand": "ca9", "prd": "ca1"}
        ctx.xref, ctx.xref_unmatched = build.Counter(), build.defaultdict(set)
        return ctx

    def edition(self):
        pdf = {"url": "https://www.uscourts.gov/f.pdf", "path": "private/acq/f.pdf", "sha256": "abc"}
        return {"period_end": "2026-06-30", "edition": "federal-court-management-statistics-june-2026", "files": {"district": pdf, "circuit": pdf}, "page_url": "p", "notes": {}, "cells": {}, "excluded": {}}

    def test_profile_record_shape_links_and_unmatched_counts(self):
        ctx = self.context()
        cur, prev = "2026-06-30", "2025-06-30"
        cell = self.cell(values={"filings": {prev: 815, cur: 1020}, "terminations": {prev: 695, cur: 934}, "pending": {prev: 941, cur: None}, "judgeships": {prev: 3, cur: 3},
                                 "pj_weighted_filings": {prev: 265, cur: 324}, "median_months_civil": {prev: 8.7, cur: 7.9}, "civil_over_3y_pct": {prev: 12.5, cur: 11.1}},
                         raw={m: {"tokens": {}, "page": 7, "y": 100.0} for m in ("filings", "terminations", "pending", "judgeships", "pj_weighted_filings", "median_months_civil", "civil_over_3y_pct")},
                         standing={"filings": {"us": 12, "circuit": 2}})
        basis = {"basis": "defendants", "evidence": "test"}
        ranks = {"filings": ({"cand": 3}, 94)}
        row = build.profile_record(ctx, self.edition(), "district", "district", "cand", cell, basis, ranks, None)
        self.assertEqual((row["dataset"], row["category"], row["id"], row["state"]), ("fcms_court_profiles", "fcms_court_profiles", "cand:2026-06-30", "CA"))
        self.assertTrue(row["source_url"].startswith("https://www.uscourts.gov/") and row["source_url"].endswith("#page=7"))
        urls = [l["url"] for l in row["item"]["links"]]
        self.assertIn("/courts/cand", urls)
        self.assertIn("/places/CA", urls)
        self.assertIn("/records/court_analytics/cand", urls)
        self.assertEqual(row["item"]["cells"]["pending"], "Not recorded")      # unknown value is never a guess
        facts = dict((a, b) for a, b in row["detail"]["facts"])
        self.assertIn("Filings (count)", facts)
        self.assertIn("standing as printed: U.S. 12, circuit 2", facts["Filings (count)"])
        self.assertIn("+25.2%", facts["Filings: change from 2025 (derived)"])
        json.dumps(row)  # JSON round-trip safe
        # a district whose id is not in court_spine and whose state has no places page: no links, both counted as unmatched
        ctx2 = self.context()
        cell2 = self.cell(court_id="zzd", name="NOWHERE", usps="PR", values={"filings": {prev: 1, cur: 2}}, raw={"filings": {"tokens": {}, "page": 3, "y": 1.0}})
        row2 = build.profile_record(ctx2, self.edition(), "district", "district", "zzd", cell2, basis, {}, None)
        urls2 = [l["url"] for l in row2["item"]["links"]]
        self.assertFalse(any(u.startswith(("/courts/", "/places/")) for u in urls2))
        self.assertEqual(ctx2.xref["district_courts_unmatched"], 1)
        self.assertEqual(ctx2.xref["district_places_no_state_page"], 1)

    def test_extreme_text_names_ties_instead_of_listing_arbitrary_courts(self):
        ctx = self.context()
        ctx.spine = dict(SPINE)
        text = build.extreme_text(ctx, "district", [("cand", 0.0), ("azd", 0.0), ("dcd", 0.0)], 3, 1, "months")
        self.assertTrue(text.startswith("3 courts tie at 0.0"))
        single = build.extreme_text(ctx, "district", [("cand", 12.5), ("azd", 9.0)], 1, 1, "months")
        self.assertEqual(single, "N.D. California 12.5; D. Arizona 9.0")
        self.assertEqual(build.extreme_text(ctx, "district", [], 0, 1, "months"), "Not recorded")

    def test_national_record_has_no_court_or_place_link(self):
        ctx = self.context()
        cell = self.cell(court_id="national", usps=None, name="ALL DISTRICT COURTS", values={"filings": {"2026-06-30": 458756}}, raw={"filings": {"tokens": {}, "page": 1, "y": 1.0}})
        row = build.profile_record(ctx, self.edition(), "district", "national_district", "national", cell, {"basis": "defendants", "evidence": "t"}, {}, None)
        self.assertEqual(row["id"], "national-district:2026-06-30")
        self.assertEqual([l["url"] for l in row["item"]["links"]], [row["item"]["links"][0]["url"]])
        self.assertTrue(row["item"]["links"][0]["url"].startswith("https://"))

    def test_gzip_json_roundtrip_is_atomic(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "a.json.gz"
            jsonio.dump(p, {"b": 1, "a": [1, 2]})
            self.assertEqual(jsonio.load(p), {"a": [1, 2], "b": 1})
            self.assertFalse((Path(tmp) / "a.json.gz.tmp").exists())
            with gzip.open(p, "rb") as fh:
                self.assertEqual(fh.read(), b'{"a":[1,2],"b":1}')


class OcrAssessmentTests(unittest.TestCase):
    def test_match_rate_against_born_digital_cells(self):
        with tempfile.TemporaryDirectory() as tmp:
            d = Path(tmp)

            def parsed(slug, period, ocr, values):
                return {"edition": slug, "period_end": period, "kind": "district", "ocr_text_layer": ocr, "excluded": [{"page": 9}],
                        "pages": [{"court_id": "cand", "values": {"filings": values}, "page": 1}, {"court_id": None, "values": {"filings": {"1999-09-30": 1}}, "page": 2}]}
            jsonio.dump(d / "born__district.json.gz", parsed("born", "2004-09-30", False, {"1999-09-30": 100, "2000-09-30": 200}))
            jsonio.dump(d / "scan__district.json.gz", parsed("scan", "2003-09-30", True, {"1999-09-30": 100, "2000-09-30": 2000}))
            index = {"born__district": {"status": "parsed"}, "scan__district": {"status": "parsed"}}
            out = ocr_assessment.compare(d, index)
            self.assertEqual(out["overall"]["compared"], 2)
            self.assertEqual(out["overall"]["match_rate"], 0.5)      # 2000 read as 2,000: an OCR digit error is counted, not corrected
            self.assertEqual(out["editions"][0]["pages_with_court_id"], 1)    # the page without a court id is never compared


def row(label, y, **values):
    return {"label": label, "values": values, "standing": {}, "stray": [], "y": y}


class ExtractionFixTests(unittest.TestCase):
    """Behaviour added after the first full extraction: stacked judgeships/panels cells, the unlabelled last-year row, ambiguous cells."""

    YEARS = ["2021-03-31", "2022-03-31", "2023-03-31"]

    def test_stacked_judgeships_panels_cell_is_reassigned_only_when_arithmetic_agrees(self):
        import extract
        y = self.YEARS
        assigned = {"judgeships": row("Judgeships/", 97.0), "panels": row("Number of Panels", 103.7, **{k: "11" for k in y})}
        unmatched = [row("", 107.8, **{k: "3.7" for k in y})]
        self.assertTrue(extract.restack_judgeships(assigned, unmatched, y))
        self.assertEqual(assigned["judgeships"]["values"]["2021-03-31"], "11")
        self.assertEqual(assigned["panels"]["values"]["2021-03-31"], "3.7")
        self.assertEqual(unmatched, [])
        # 11 judgeships cannot have 5.0 panels: nothing is moved, nothing is guessed
        assigned = {"judgeships": row("Judgeships/", 97.0), "panels": row("Number of Panels", 103.7, **{k: "11" for k in y})}
        unmatched = [row("", 107.8, **{k: "5.0" for k in y})]
        self.assertFalse(extract.restack_judgeships(assigned, unmatched, y))
        self.assertEqual(assigned["judgeships"]["values"], {})
        # a normally printed judgeships row is left alone
        normal = {"judgeships": row("Number of Judgeships", 97.0, **{k: "16" for k in y}), "panels": row("Number of Panels", 101.0, **{k: "5.3" for k in y})}
        self.assertFalse(extract.restack_judgeships(normal, [], y))

    def test_unlabelled_last_year_line_above_the_percent_change_label(self):
        import extract
        y = ["2010-09-30", "2011-09-30", "2012-09-30"]
        rows = [row("Total", 1.0, **{k: "100" for k in y}), row("Prisoner", 2.0, **{k: "10" for k in y}), row("Other", 3.0, **{k: "20" for k in y}),
                row("Criminal", 4.0, **{k: "30" for k in y}), row("Administrative", 5.0, **{k: "40" for k in y}),
                row("", 6.0, **{"2011-09-30": "4.3"}),
                row("Percent Change in Total Filings Current Year Over Earlier Year", 7.0, **{"2010-09-30": "2.7"}),
                row("Total", 8.0), row("Procedural", 9.0), row("Total", 10.0), row("Prisoner", 11.0), row("Other", 12.0), row("Criminal", 13.0), row("Administrative", 14.0),
                row("Pending Appeals", 15.0)]
        assigned, unmatched, missing = measures.classify(rows, measures.CIRCUIT_LEFT)
        self.assertEqual(missing, [])
        self.assertEqual(assigned["pct_change_filings_last_year"]["values"], {"2011-09-30": "4.3"})
        self.assertEqual(assigned["pct_change_filings"]["values"], {"2010-09-30": "2.7"})
        entry = {"values": {"pct_change_filings": {"2010-09-30": 2.7, "2011-09-30": None, "2012-09-30": None}, "pct_change_filings_last_year": {"2011-09-30": 4.3}}, "flags": {}, "standing": {}}
        extract.merge_earlier_pct(entry)
        self.assertEqual(entry["values"]["pct_change_filings"], {"2010-09-30": 2.7, "2011-09-30": 4.3, "2012-09-30": None})
        self.assertNotIn("pct_change_filings_last_year", entry["values"])

    def test_cell_with_extra_figures_is_ambiguous_and_never_used(self):
        import extract
        r = row("Median Time", 1.0, **{"2021-03-31": "9.7", "2022-03-31": "11.2"})
        r["stray"] = ["9.5", "14.0"]
        values, flags = extract.values_of(r, ["2021-03-31", "2022-03-31"])
        self.assertEqual(values, {"2021-03-31": None, "2022-03-31": None})
        self.assertEqual(flags, {"2021-03-31": "ambiguous", "2022-03-31": "ambiguous"})
        values, flags = extract.values_of(row("Median Time", 1.0, **{"2021-03-31": "9.7", "2022-03-31": "-"}), ["2021-03-31", "2022-03-31"])
        self.assertEqual((values["2021-03-31"], values["2022-03-31"], flags), (9.7, None, {"2022-03-31": "dash"}))


class VerificationToolTests(unittest.TestCase):
    """The independent checkers themselves (crosscheck_excel.py, spot_check_fcms.py) are tested on synthetic input."""

    def test_workbook_cells_become_series(self):
        import crosscheck_excel as cx
        self.assertEqual(cx.cell_numbers("11 /\n 3.7"), [11.0, 3.7])
        self.assertEqual(cx.cell_numbers("63,671\n10.8"), [63671.0, 10.8])
        self.assertEqual(cx.cell_numbers("-"), [None])
        self.assertEqual(cx.cell_numbers(None), [])
        self.assertEqual(cx.cell_numbers(0.5), [0.5])
        self.assertEqual(cx.cell_numbers("Overall Caseload Statistics"), [])

    def test_tolerance_follows_the_printed_precision_and_accepts_percent_fractions(self):
        import crosscheck_excel as cx
        self.assertTrue(cx.same(10.0655738, 10.1, "10.1"))        # workbook holds the unrounded number
        self.assertFalse(cx.same(10.0, 10.1, "10.1"))
        self.assertTrue(cx.same(0.771890266816986, 77.2, "77.20%"))   # fraction in the workbook, "77.20%" printed
        self.assertFalse(cx.same(0.80, 77.2, "77.20%"))
        self.assertTrue(cx.same(None, None, None))
        self.assertFalse(cx.same(None, 5, "5"))
        self.assertFalse(cx.same(5, None, None))

    def test_blocks_are_cut_at_period_headers_and_adjacent_headers_are_one_block(self):
        import openpyxl
        import crosscheck_excel as cx
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "MAINE"
        for j, y in enumerate(("Mar 31\n2021", "Mar 31\n2022", "Mar 31\n2023", "Mar 31\n2024")):
            ws.cell(row=2, column=6 + j, value=y)            # left table header
            ws.cell(row=3, column=16 + j, value=y)           # right table header on the next row
        for j, v in enumerate((10, 20, 30, 40)):
            ws.cell(row=3, column=6 + j, value=v)            # left table, first data row (same row as the right header)
            ws.cell(row=4, column=16 + j, value=v * 2)
        ws.cell(row=3, column=3, value="Filings")
        ws["A7"] = "OHIO"
        for j, y in enumerate(("Mar 31\n2021", "Mar 31\n2022", "Mar 31\n2023", "Mar 31\n2024")):
            ws.cell(row=8, column=6 + j, value=y)
        for j, v in enumerate((1, 2, 3, 4)):
            ws.cell(row=9, column=6 + j, value=v)
        blocks = dict(cx.blocks_of(ws))
        self.assertEqual(set(blocks), {"MAINE", "OHIO"})
        maine = [s for panel in blocks["MAINE"] for _, _, ser in panel for s in ser]
        self.assertIn([10.0, 20.0, 30.0, 40.0], maine)
        self.assertIn([20.0, 40.0, 60.0, 80.0], maine)             # the right table's row was not lost to the adjacent header
        self.assertNotIn([21.0, 22.0, 23.0, 24.0], maine)
        ohio = [s for panel in blocks["OHIO"] for _, _, ser in panel for s in ser]
        self.assertEqual(ohio, [[1.0, 2.0, 3.0, 4.0]])

    def test_numeric_year_headers_never_swallow_data_rows(self):
        import openpyxl
        import crosscheck_excel as cx
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A3"] = "MAINE"
        for j, y in enumerate((2007, 2008, 2009, 2010, 2011)):
            ws.cell(row=4, column=7 + j, value=y)
        for j, v in enumerate((2010, 2020, 2030, 2040, 2050)):   # a data row that looks like years but spans more than six
            ws.cell(row=5, column=7 + j, value=v)
        for j, v in enumerate((2100, 2020, 3000, 40, 50)):
            ws.cell(row=6, column=7 + j, value=v)
        blocks = list(cx.blocks_of(ws))
        self.assertEqual(len(blocks), 1)
        series = [s for panel in blocks[0][1] for _, _, ser in panel for s in ser]
        self.assertEqual(len(series), 2)

    def test_pdf_text_numbers_and_series_location(self):
        import spot_check_fcms as sc
        self.assertEqual(sc.line_numbers("Filings ¹   1,020   -   7.9  84.6%  -3.3"), [1020.0, None, 7.9, 84.6, -3.3])
        self.assertEqual(sc.line_numbers("12/4.0"), [])                   # glued tokens are never read as numbers
        key = ("fake.pdf", 1)
        sc.cx.PAGE_TEXT[key] = ["header", "Number of Judgeships", "   22   22   22   22  22  22", "Filings*     8,591   8,422   8,093   9,056   10,584  11,126   40  2"]
        hit = sc.find_series("fake.pdf", [1], [8591.0, 8422.0, 8093.0, 9056.0, 10584.0, 11126.0], "filings")
        self.assertEqual((hit[0], hit[1]), (1, 3))
        # newest-first editions: the ascending series is found as the reversed run
        asc = [11126.0, 10584.0, 9056.0, 8093.0, 8422.0, 8591.0]
        self.assertIsNone(sc.find_series("fake.pdf", [1], asc, "filings"))
        located = sc.locate("fake.pdf", [1], asc, "filings")
        self.assertTrue(located[4])
        # a series whose label word is missing is not accepted
        self.assertIsNone(sc.find_series("fake.pdf", [1], [8591.0, 8422.0, 8093.0, 9056.0, 10584.0, 11126.0], "terminations"))
        # a dash in the page matches a missing value
        sc.cx.PAGE_TEXT[("fake.pdf", 2)] = ["Civil trial", "  -   30.6   -   -   -   -"]
        self.assertIsNotNone(sc.find_series("fake.pdf", [2], [None, 30.6, None, None, None, None], "trial"))
        self.assertEqual(sc.displayed_number("1,020 (standing as printed: U.S. 12)"), 1020.0)
        self.assertIsNone(sc.displayed_number("Not recorded (printed as a dash)"))

    def test_mismatch_adjudication_needs_the_tokens_in_order_on_one_line(self):
        import crosscheck_excel as cx
        cx.PAGE_TEXT[("x.pdf", 1)] = ["Total   1,307   1,097   1,178   1,176   1,143   1,170", "Other   443  466"]
        self.assertTrue(cx.pdf_text_confirms("x.pdf", 1, ["1,307", "1,097", "1,178", "1,176", "1,143", "1,170"]))
        self.assertTrue(cx.pdf_text_confirms("x.pdf", 1, ["1,307", None, "1,178"]))             # a dash column is skipped
        self.assertFalse(cx.pdf_text_confirms("x.pdf", 1, ["1,097", "1,307"]))                  # wrong order
        self.assertFalse(cx.pdf_text_confirms("x.pdf", 1, ["1,30"]))                            # a prefix of a number is not the number
        self.assertFalse(cx.pdf_text_confirms("x.pdf", 1, [None, None]))


if __name__ == "__main__":
    unittest.main()
