"""Offline tests for the uscourts_tables parsers (synthetic workbooks, no network).

Run: cd scripts/enrichment/uscourts_tables && python3 -B -m unittest -v test_uscourts_tables
"""
import io
import unittest

import openpyxl

import entities as E
import ingest as G
import utx_parse as P


def district_workbook(period="September 30, 2025", dash=True, text_zero=True):
    """A small workbook shaped like Judicial Business Table T-1 (merged headings, circuit roll-ups, notes)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Table T-1."
    ws["A2"] = f"Table T-1._x000D_\nU.S. District Courts––Civil and Criminal Trials Completed, by District,_x000D_\nDuring the 12-Month Period Ending {period}"
    ws.merge_cells("A2:H2")
    ws["A3"] = "Circuit and District"
    ws.merge_cells("A3:A4")
    ws["B3"] = "Total\nAll Trials"
    ws.merge_cells("B3:B4")
    ws["C3"] = "Civil ¹"
    ws.merge_cells("C3:E3")
    ws["F3"] = "Criminal"
    ws.merge_cells("F3:H3")
    for col, text in zip("CDEFGH", ["Total", "Nonjury", "Jury", "Total", "Nonjury", "Jury"]):
        ws[f"{col}4"] = text
    rows = [
        ("       Total", [10, 6, 4, 2, 4, 3, 1]),
        ("DC", [3, 2, 2, 0, 1, 1, 0]),
        ("       1st", [7, 4, 2, 2, 3, 2, 1]),
        ("ME", [3, 2, 1, 1, 1, 1, 0]),
        ("NY,S", [4, 2, 1, 1, 2, 1, 1]),
    ]
    r = 5
    for label, values in rows:
        ws.cell(r, 1, label)
        for c, v in enumerate(values, 2):
            ws.cell(r, c, v)
        r += 1
    if dash:
        ws.cell(r - 1, 5, " - ")
    if text_zero:
        ws.cell(r - 2, 8, "0")
    ws.cell(r + 1, 1, "NOTE: This table excludes sentencing hearings._x000D_\n¹ Includes trials of miscellaneous cases.")
    ws.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=8)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class ParseTests(unittest.TestCase):
    def setUp(self):
        sheets, problems = P.parse_workbook(district_workbook())
        self.assertEqual(problems, [])
        self.assertEqual(len(sheets), 1)
        self.sheet = sheets[0]

    def test_title_period_and_table_id(self):
        s = self.sheet
        self.assertEqual(s["table_id"], "T-1")
        self.assertIn("Civil and Criminal Trials Completed, by District", s["title_raw"])
        self.assertEqual(s["period_end_from_title"], "2025-09-30")
        self.assertTrue(s["period_phrase"].startswith("During the 12-Month Period Ending"))

    def test_column_paths_keep_published_headings_without_footnote_marks(self):
        paths = [c["path"] for c in self.sheet["columns"]]
        self.assertEqual(paths[0], ["Total All Trials"])
        self.assertEqual(paths[1], ["Civil", "Total"])
        self.assertEqual(paths[5], ["Criminal", "Nonjury"] if False else ["Criminal", "Nonjury"])
        self.assertEqual(len(paths), 7)

    def test_rows_and_labels(self):
        labels = [r["label"] for r in self.sheet["rows"]]
        self.assertEqual(labels, ["Total", "DC", "1st", "ME", "NY,S"])
        total = self.sheet["rows"][0]
        self.assertEqual(total["values"][0], 10)
        self.assertEqual(total["label_raw"].strip(), "Total")

    def test_dash_is_not_zero_and_text_number_is_flagged(self):
        last = self.sheet["rows"][-1]
        self.assertIsNone(last["values"][3])
        self.assertEqual(last["flags"][3], "dash")
        dc_zero = self.sheet["rows"][-2]
        self.assertEqual(dc_zero["values"][-1], 0)
        self.assertEqual(dc_zero["flags"][-1], "textnum")

    def test_notes_split_into_published_lines(self):
        notes = self.sheet["notes"]
        self.assertEqual(notes[0], "NOTE: This table excludes sentencing hearings.")
        self.assertEqual(notes[1], "¹ Includes trials of miscellaneous cases.")

    def test_footnote_marks_on_headings_and_notes_are_kept(self):
        s = self.sheet
        civil = [c for c in s["columns"] if c["path"][:1] == ["Civil"]]
        self.assertEqual(civil[0]["marks"], ["¹"])
        self.assertEqual(s["note_marks"]["¹"], "¹ Includes trials of miscellaneous cases.")
        self.assertEqual(s["columns"][0]["marks"], [])

    def test_footnote_mark_on_a_row_label_is_kept_and_stripped_from_the_label(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "Table C. U.S. District Courts During the 12-Month Period Ending September 30, 2025"
        ws.append(["Circuit and District", "Filed", "Pending"])
        ws.append(["Total", 5, 4])
        ws.append(["WV,S²", 1, 2])
        ws.append([None])
        ws.append(["NOTE: x"])
        ws.append(["² The higher number of pending cases is the result of a transfer."])
        buf = io.BytesIO()
        wb.save(buf)
        sheets, _ = P.parse_workbook(buf.getvalue())
        row = sheets[0]["rows"][1]
        self.assertEqual(row["label"], "WV,S")
        self.assertEqual(row["marks"], ["²"])
        self.assertIn("transfer", sheets[0]["note_marks"]["²"])

    def test_labelled_rows_under_the_footnote_block_are_notes_not_data(self):
        # Bankruptcy F-2 (2016-2018) prints "Chapter 9: MO, E 1 ..." fragments, with district labels, under its footnotes
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "Table F-2. U.S. Bankruptcy Courts During the 12-Month Period Ending December 31, 2018"
        ws.append(["Circuit and District", "Total", "Chapter 7"])
        ws.append(["Total", 5, 4])
        ws.append(["ME", 2, 2])
        ws.append(["NY, S", 3, 2])
        ws.append([None])
        ws.append(["¹ The nature of debt is business if the debtor is a corporation."])
        ws.append(["Note: These figures include the following cases not reflected elsewhere:"])
        ws.append(["Chapter 9"])
        ws.append(["MO, E", 1, "085", "NY, E", 1, "022"])
        buf = io.BytesIO()
        wb.save(buf)
        sheets, _ = P.parse_workbook(buf.getvalue())
        self.assertEqual([r["label"] for r in sheets[0]["rows"]], ["Total", "ME", "NY, S"])
        self.assertEqual(len(sheets[0]["columns"]), 2)
        self.assertIn("MO, E  1  085  NY, E  1  022", sheets[0]["notes"])
        self.assertIn("Chapter 9", sheets[0]["notes"])

    def test_a_sort_order_helper_column_is_not_a_statistic(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "Table F-2. U.S. Bankruptcy Courts During the 12-Month Period Ending March 31, 2018"
        ws.append(["Circuit and District", "Total Filings", "Total Chapter 7", "Sort Order"])
        ws.append(["Total", 5, 4, "0.0"])
        ws.append(["DC", 2, 2, "00 "])
        ws.append(["ME", 3, 2, "010"])
        buf = io.BytesIO()
        wb.save(buf)
        sheets, _ = P.parse_workbook(buf.getvalue())
        self.assertEqual([c["path"] for c in sheets[0]["columns"]], [["Total Filings"], ["Total Chapter 7"]])
        self.assertEqual(sheets[0]["rows"][1]["values"], [2, 2])

    def test_year_header_row_is_not_data(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "Table C. U.S. District Courts During the 12-Month Period Ending September 30, 2025"
        ws.append(["Circuit and District", "Filed", None, "Terminated", None])
        ws.append([None, 2025, 2024, 2025, 2024])
        ws.append(["Total", 5, 4, 3, 2])
        ws.append(["ME", 1, 1, 1, 1])
        buf = io.BytesIO()
        wb.save(buf)
        sheets, _ = P.parse_workbook(buf.getvalue())
        self.assertEqual([r["label"] for r in sheets[0]["rows"]], ["Total", "ME"])
        self.assertEqual(sheets[0]["columns"][0]["path"], ["Filed", "2025"])


class ValueTests(unittest.TestCase):
    def test_to_value(self):
        self.assertEqual(P.to_value(12), (12, None))
        self.assertEqual(P.to_value(1.5), (1.5, None))
        self.assertEqual(P.to_value(" - "), (None, "dash"))
        self.assertEqual(P.to_value("1,234"), (1234, "textnum"))
        self.assertEqual(P.to_value("0"), (0, "textnum"))
        self.assertEqual(P.to_value(None), (None, "blank"))
        self.assertEqual(P.to_value("NA"), (None, "na"))
        value, flag = P.to_value("see note")
        self.assertIsNone(value)
        self.assertTrue(flag.startswith("text:"))

    def test_period_phrase(self):
        phrase, end = P.parse_period_phrase("Cases Commenced, During the 12-Month Period Ending June 30, 2026")
        self.assertEqual(end, "2026-06-30")
        phrase, end = P.parse_period_phrase("During the 12-Month Periods Ending September 30, 2023 Through 2025")
        self.assertEqual(end, "2025-09-30")
        phrase, end = P.parse_period_phrase("Cases Commenced, Terminated, and Pending During the 12-Month Periods Ending March 31, 2025 and 2026")
        self.assertEqual(end, "2026-03-31")
        self.assertTrue(phrase.endswith("2025 and 2026"))
        # a period with no day ends on the month's last day; a later "as of" extraction date is not the period
        phrase, end = P.parse_period_phrase("Civil Consent Cases During the 12-Month Period Ending September, 2020 as of November 18, 2020")
        self.assertEqual(end, "2020-09-30")
        phrase, end = P.parse_period_phrase("Civil Cases Pending, by District and Length of Time Pending, as of September 30, 2025")
        self.assertEqual(end, "2025-09-30")
        self.assertTrue(phrase.lower().startswith("as of"))
        phrase, end = P.parse_period_phrase("During the 12-Month Period Ending February, 2024")
        self.assertEqual(end, "2024-02-29")
        self.assertEqual(P.parse_period_phrase("No period here"), (None, None))


class EntityTests(unittest.TestCase):
    def test_district_ids(self):
        self.assertEqual(E.district_id("NY,S"), "nysd")
        self.assertEqual(E.district_id("CA,N"), "cand")
        self.assertEqual(E.district_id("ME"), "med")
        self.assertEqual(E.district_id("DC"), "dcd")
        self.assertEqual(E.district_id("PR"), "prd")
        self.assertEqual(E.district_id("GU"), "gud")
        self.assertEqual(E.district_id("NMI"), "nmid")
        self.assertEqual(E.district_id("WV,S2"), "wvsd")
        self.assertEqual(E.district_id("NY, S    "), "nysd")  # older editions print a space after the comma
        self.assertEqual(E.district_id("TX,S", kind="b"), "txsb")
        self.assertEqual(E.district_id("AR", kind="b"), "arb")

    def test_non_districts_are_none(self):
        for label in ("Total", "1st", "Contract Actions", "", "Circuit and District", "NOTE: x"):
            self.assertIsNone(E.district_id(label), label)

    def test_circuit_ids(self):
        self.assertEqual(E.circuit_id("1st"), "ca1")
        self.assertEqual(E.circuit_id("     11th"), "ca11")
        self.assertEqual(E.circuit_id("DC"), "cadc")
        self.assertEqual(E.circuit_id("Fed."), "cafc")
        self.assertIsNone(E.circuit_id("ME"))

    def test_national(self):
        self.assertTrue(E.is_national("   Total"))
        self.assertFalse(E.is_national("1st"))


class AcceptanceTests(unittest.TestCase):
    @staticmethod
    def norm(columns, values):
        return {"rows": [{"flags": [None] * len(values), "values": values}], "printed_table_id": "C", "table": "C", "title": "t", "title_listing": "t", "scope": "district",
                "problems": [], "columns": columns}

    def test_a_clean_table_is_accepted(self):
        norm = self.norm([{"key": "a", "path": ["A"], "index": 1}, {"key": "b", "path": ["B"], "index": 2}], [1, 2])
        self.assertEqual(G.accept(norm), [])

    def test_two_columns_with_one_heading_path_are_rejected(self):
        norm = self.norm([{"key": "total", "path": ["Total"], "index": 1}, {"key": "total~2", "path": ["Total"], "index": 2}], [1, 2])
        self.assertTrue(any("ambiguous column headings" in r for r in G.accept(norm)))

    def test_a_category_table_that_repeats_a_label_with_different_figures_is_rejected(self):
        norm = self.norm([{"key": "a", "path": ["A"], "index": 1}], [1])
        norm["scope"] = "category"
        norm["rows"] = [{"entity_id": "franchise", "flags": [None], "values": [1]}, {"entity_id": "other", "flags": [None], "values": [5]},
                        {"entity_id": "franchise", "flags": [None], "values": [2]}]
        self.assertTrue(any("ambiguous row labels" in r for r in G.accept(norm)))
        norm["rows"][2]["values"] = [1]  # the same figures twice are one row
        self.assertEqual(G.accept(norm), [])

    def test_a_column_with_figures_but_no_heading_is_rejected(self):
        norm = self.norm([{"key": "a", "path": ["A"], "index": 1}, {"key": "col", "path": [], "index": 2}], [1, 2])
        self.assertTrue(any("no heading" in r for r in G.accept(norm)))
        # an empty headingless column (no figures at all) is harmless
        self.assertEqual(G.accept(self.norm([{"key": "a", "path": ["A"], "index": 1}, {"key": "col", "path": [], "index": 2}], [1, None])), [])


if __name__ == "__main__":
    unittest.main()
