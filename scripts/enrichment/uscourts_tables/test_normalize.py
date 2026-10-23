"""Offline tests for normalize.py and checks.py (synthetic workbooks shaped like the published district tables)."""
import io
import json
import unittest
from pathlib import Path

import openpyxl

import checks as C
import normalize as N
import utx_parse as P

CIRCUITS = [
    ("DC", ["DC"]),
    ("1st", ["ME", "MA", "NH", "RI", "PR"]),
    ("2nd", ["CT", "NY,N", "NY,E", "NY,S", "NY,W", "VT"]),
    ("3rd", ["DE", "NJ", "PA,E", "PA,M", "PA,W", "VI"]),
    ("4th", ["MD", "NC,E", "NC,M", "NC,W", "SC", "VA,E", "VA,W", "WV,N", "WV,S"]),
    ("5th", ["LA,E", "LA,M", "LA,W", "MS,N", "MS,S", "TX,N", "TX,E", "TX,S", "TX,W"]),
    ("6th", ["KY,E", "KY,W", "MI,E", "MI,W", "OH,N", "OH,S", "TN,E", "TN,M", "TN,W"]),
    ("7th", ["IL,N", "IL,C", "IL,S", "IN,N", "IN,S", "WI,E", "WI,W"]),
    ("8th", ["AR,E", "AR,W", "IA,N", "IA,S", "MN", "MO,E", "MO,W", "NE", "ND", "SD"]),
    ("9th", ["AK", "AZ", "CA,N", "CA,E", "CA,C", "CA,S", "GU", "HI", "ID", "MT", "NV", "NMI", "OR", "WA,E", "WA,W"]),
    ("10th", ["CO", "KS", "NM", "OK,N", "OK,E", "OK,W", "UT", "WY"]),
    ("11th", ["AL,N", "AL,M", "AL,S", "FL,N", "FL,M", "FL,S", "GA,N", "GA,M", "GA,S"]),
]
IDS = json.loads((Path(__file__).resolve().parents[3] / "private" / "build" / "uscourts_tables" / "lookups" / "court_ids.json").read_text()) if (Path(__file__).resolve().parents[3] / "private" / "build" / "uscourts_tables" / "lookups" / "court_ids.json").exists() else None


def table_workbook(break_sum=False, table="C-4A"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Table {table}"
    ws["A1"] = f"Table {table}. U.S. District Courts––Civil Cases Terminated, by District and Action Taken, During the 12-Month Period Ending September 30, 2025"
    ws["A2"], ws["B2"], ws["C2"], ws["D2"] = "Circuit and District", "Total", "No Court Action", "Court Action"
    ws["A3"], ws["D3"], ws["E3"] = None, "Total", "Percent Reaching Trial"
    rows, total = [], [0, 0, 0]
    n = 0
    for circuit, districts in CIRCUITS:
        crow = [0, 0, 0]
        drows = []
        for d in districts:
            n += 1
            vals = [10 + n, 4 + n % 3, 6 + n - n % 3]
            drows.append((d, vals))
            crow = [a + b for a, b in zip(crow, vals)]
        total = [a + b for a, b in zip(total, crow)]
        if circuit != "DC":
            rows.append((circuit, crow))
        rows.extend(drows)
    if break_sum:
        rows[1][1][0] += 1  # 1st circuit total no longer equals the sum of its districts
    ws.append([])
    r = 4
    ws.cell(r, 1, "Total")
    for c, v in enumerate(total, 2):
        ws.cell(r, c, v)
    ws.cell(r, 5, 0.41)
    for label, vals in rows:
        r += 1
        ws.cell(r, 1, label)
        for c, v in enumerate(vals, 2):
            ws.cell(r, c, v)
        ws.cell(r, 5, 1.5)
    ws.cell(r + 2, 1, "NOTE: Land condemnation cases omitted.")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def normalized(break_sum=False, table="C-4A", period="2025-09-30"):
    sheets, problems = P.parse_workbook(table_workbook(break_sum, table))
    assert not problems, problems
    entry = {"pub": "JB", "table": table, "period_end": period, "title": "listing title", "page": "https://example.test/p"}
    return N.normalize(sheets[0], entry)


class NormalizeTests(unittest.TestCase):
    def test_scope_and_entities(self):
        n = normalized()
        self.assertEqual(n["scope"], "district")
        kinds = {}
        for r in n["rows"]:
            kinds[r["entity"]] = kinds.get(r["entity"], 0) + 1
        self.assertEqual(kinds["national"], 1)
        self.assertEqual(kinds["district"], 94)
        self.assertEqual(kinds["circuit"], 11)  # DC is the district row; the other 11 circuits print a roll-up row
        ids = {r["entity_id"] for r in n["rows"] if r["entity"] == "district"}
        self.assertIn("nysd", ids)
        self.assertIn("nmid", ids)
        self.assertIn("gud", ids)
        if IDS:
            self.assertEqual(ids, set(IDS["districts"]))

    def test_title_period_and_months(self):
        n = normalized()
        self.assertEqual(n["period_end"], "2025-09-30")
        self.assertEqual(n["months"], 12)
        self.assertTrue(n["title"].startswith("U.S. District Courts––Civil Cases Terminated"))
        self.assertNotIn("Table C-4A", n["title"])
        self.assertNotIn("Period Ending", n["title"])
        self.assertEqual(n["printed_table_id"], "C-4A")

    def test_column_keys_and_paths(self):
        n = normalized()
        keys = [c["key"] for c in n["columns"]]
        self.assertEqual(keys[0], "total")
        self.assertIn("court_action/total", keys)
        labels = [c["label"] for c in n["columns"]]
        self.assertIn("Court Action / Total", labels)

    def test_relative_year_columns(self):
        self.assertEqual(N.column_key(["Filed", "2025"], 2025), "filed/y0")
        self.assertEqual(N.column_key(["Filed", "2024"], 2025), "filed/y-1")
        self.assertEqual(N.column_key(["Filed", "FY2023"], 2025), "filed/y-2")
        self.assertTrue(N.is_prior_year_column(["Filed", "2024"], 2025))
        self.assertFalse(N.is_prior_year_column(["Filed", "2025"], 2025))
        self.assertFalse(N.is_prior_year_column(["Percent Change", "2024"], 2025))

    def test_dates_inside_headings_are_relative(self):
        a = N.column_key(["Not Terminated as of June 30, 2024", "Felony"], 2025)
        b = N.column_key(["Not Terminated as of June 30, 2025", "Felony"], 2026)
        self.assertEqual(a, b)
        self.assertEqual(a, "not_terminated_as_of_june_30_ym1/felony")
        self.assertNotEqual(a, N.column_key(["Not Terminated as of June 30, 2025", "Felony"], 2025))
        self.assertEqual(N.column_key(["Filings 2024 and 2025"], 2025), "filings_ym1_and_y0")

    def test_renamed_top_level_headings_are_one_column(self):
        self.assertEqual(N.column_key(["Filings", "2021"], 2021), N.column_key(["Filed", "2022"], 2022))
        self.assertEqual(N.column_key(["Commenced", "2024"], 2024), "filed/y0")
        self.assertEqual(N.column_key(["Terminations", "2021"], 2021), N.column_key(["Terminated", "2026"], 2026))
        # only the exact top-level heading is merged: a lower-level word or a longer heading is its own column
        self.assertNotEqual(N.column_key(["Total", "Filings"], 2025), N.column_key(["Total", "Filed"], 2025))
        self.assertEqual(N.column_key(["Total Filings"], 2025), "total_filings")

    def test_table_slug_keeps_variants_apart(self):
        self.assertEqual(N.table_slug("C-4A"), "c-4a")
        self.assertEqual(N.table_slug("D Cases"), "d-cases")
        self.assertEqual(N.table_slug("F-2 (One Month)"), "f-2-one-month")
        self.assertNotEqual(N.table_slug("D-1"), N.table_slug("D-1 Cases"))


def county_workbook():
    """A workbook shaped like Bankruptcy Filings F-5A: district rows, each followed by county rows with a FIPS County Code column."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Table F-5A"
    ws["A2"] = "Report F-5A.\nU.S. Bankruptcy Courts—Business and Nonbusiness Bankruptcy Cases Commenced, by County and Chapter of the Bankruptcy Code,\nDuring the 12-Month Period Ending June 30, 2026"
    ws.merge_cells("A2:D2")
    ws.append([])
    ws.append(["Circuit,\nDistrict, and\nCounty", "County\nCode ¹", "Total\nAll\nChapters", "Total\nChapter\n7"])
    ws.append(["Total", None, 100, 60])
    n = 0
    for circuit, districts in CIRCUITS:
        if circuit != "DC":
            ws.append([f"     {circuit}", None, 1, 1])
        for d in districts:
            n += 1
            ws.append([d, None, 5, 3])
            code = "01001" if n == 1 else f"{n:02d}{n:03d}"
            ws.append([f"COUNTY {n}", code, 4, 2])
            if n == 2:
                ws.append(["OUTSIDE COUNTY", "12113*", 1, 1])
                ws.append(["NUMERIC CODE", 1003, 1, "-"])
    ws.append([None])
    ws.append(["NOTE: x"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class CountyTests(unittest.TestCase):
    def setUp(self):
        sheets, problems = P.parse_workbook(county_workbook())
        self.assertEqual(problems, [])
        entry = {"pub": "BF", "table": "F-5A", "period_end": "2026-06-30", "title": "listing", "page": "https://example.test/p"}
        self.norm = N.normalize(sheets[0], entry)

    def test_county_table_is_its_own_scope_and_the_code_is_not_a_series(self):
        n = self.norm
        self.assertEqual(n["scope"], "county")
        self.assertTrue([c for c in n["columns"] if c["label"].startswith("County Code")][0]["prior"])
        self.assertEqual([c["label"] for c in n["columns"] if not c["prior"]], ["Total All Chapters", "Total Chapter 7"])

    def test_counties_carry_their_fips_code_and_district(self):
        counties = [r for r in self.norm["rows"] if r["entity"] == "county"]
        self.assertEqual(len(counties), 94 + 2)
        first = counties[0]
        self.assertEqual(first["county_code"], "01001")  # the leading zero is kept
        self.assertEqual(first["district"], "dcd")
        self.assertEqual(first["entity_id"], "01001:dcd")
        outside = next(r for r in counties if r["label"] == "OUTSIDE COUNTY")
        self.assertEqual(outside["county_code"], "12113")
        self.assertTrue(outside["outside_district"])
        numeric = next(r for r in counties if r["label"] == "NUMERIC CODE")
        self.assertEqual(numeric["county_code"], "01003")
        self.assertIsNone(numeric["values"][2])  # the dash stays a dash
        self.assertEqual(numeric["flags"][2], "dash")

    def test_district_rows_are_not_counties(self):
        kinds = {}
        for r in self.norm["rows"]:
            kinds[r["entity"]] = kinds.get(r["entity"], 0) + 1
        self.assertEqual(kinds["district"], 94)
        self.assertEqual(kinds["circuit"], 11)


class LegacyLayoutTests(unittest.TestCase):
    def sheet(self, labels):
        rows = [{"label_raw": label, "label": label, "indent": 0, "section": False, "values": [1], "flags": [None], "row": i + 1, "marks": [], "cell_marks": {}}
                for i, label in enumerate(labels)]
        return {"sheet": "Sheet1", "table_id": None, "title_raw": "TABLE F-2. U.S. BANKRUPTCY COURTS DURING THE 12 MONTHS ENDED SEPTEMBER 30, 2005",
                "period_phrase": None, "period_end_from_title": None, "label_header": [], "columns": [{"index": 1, "path": ["TOTAL"], "fmt": None, "marks": []}],
                "rows": rows, "notes": [], "note_marks": {}, "header_rows": 1}

    ENTRY = {"pub": "BF", "table": "F-2", "period_end": "2005-09-30", "title": "x"}

    def test_dot_leader_printouts_are_flagged_not_parsed(self):
        labels = [f"{code}......" for code in ["DC", "ME", "MA", "NH", "RI", "PR", "CT", "VT", "DE", "NJ"] * 5]
        norm = N.normalize(self.sheet(labels), self.ENTRY)
        self.assertEqual(norm["scope"], "category")
        self.assertTrue(any("legacy printout layout" in p for p in norm["problems"]))

    def test_a_county_table_without_district_rows_is_flagged(self):
        sheet = self.sheet(["RIVERSIDE", "FAIRFIELD"])
        sheet["title_raw"] = "Table F-5A. U.S. Bankruptcy Courts Business and Nonbusiness Bankruptcy County Cases Commenced, During the 12-Month Period Ending June 30, 2015"
        norm = N.normalize(sheet, {**self.ENTRY, "table": "F-5A"})
        self.assertTrue(any("not grouped under their reporting district" in p for p in norm["problems"]))

    def test_children_of_a_row_with_figures_keep_their_parent_in_the_key(self):
        def row(label, indent, value):
            return {"label_raw": label, "label": label.strip(), "indent": indent, "section": False, "values": [value], "flags": [None], "row": 1, "marks": [], "cell_marks": {}}

        sheet = self.sheet([])
        sheet["rows"] = [row("Total Cases", 5, 30), row("Contract Actions, Total", 0, 10), row("     Marine", 5, 1), row("     Other", 5, 9),
                         row("Tort Actions, Total", 0, 20), row("     Marine", 5, 2), row("     Other", 5, 18)]
        sheet["title_raw"] = "Table C-2. U.S. District Courts Civil Cases Commenced, by Nature of Suit, During the 12-Month Period Ending September 30, 2025"
        norm = N.normalize(sheet, {**self.ENTRY, "table": "C-2"})
        keys = [r["entity_id"] for r in norm["rows"]]
        self.assertEqual(len(keys), len(set(keys)))
        self.assertIn("contract_actions_total/marine", keys)
        self.assertIn("tort_actions_total/marine", keys)
        self.assertIn("total_cases", keys)

    def test_a_short_category_table_with_a_trailing_ellipsis_is_not_flagged(self):
        norm = N.normalize(self.sheet(["Contract Actions..", "Real Property.."]), self.ENTRY)
        self.assertEqual(norm["problems"], [])


class CheckTests(unittest.TestCase):
    def test_consistent_table_has_no_rollup_failures(self):
        n = normalized()
        self.assertEqual(C.rollup_failures(n), [])

    def test_a_broken_circuit_total_is_reported_not_repaired(self):
        n = normalized(break_sum=True)
        fails = C.rollup_failures(n)
        self.assertTrue(fails)
        ents = {f[1] for f in fails}
        self.assertIn("ca1", ents)
        # the national total no longer equals the sum of the circuits either
        self.assertIn("national", ents)

    def test_percent_columns_are_not_summed(self):
        n = normalized()
        cols = [n["columns"][i]["key"] for i in C.additive_columns(n)]
        self.assertNotIn("court_action/percent_reaching_trial", cols)
        self.assertIn("total", cols)


if __name__ == "__main__":
    unittest.main()
