"""Generic parser for U.S. Courts statistical-table workbooks (Judicial Business, Statistical Tables for the Federal
Judiciary, Federal Judicial Caseload Statistics, Bankruptcy Filings, Wiretap, BAPCPA).

Every workbook is one published table: a title block ("Table C-4A. ... During the 12-Month Period Ending September 30,
2025"), merged multi-row column headings, a label column ("Circuit and District", "Nature of Suit", "Offense" ...),
data rows (with a "Total" row, circuit roll-up rows and district rows, or category rows) and "NOTE" / footnote rows.

The parser keeps the published words: row labels (raw and cleaned), full column heading paths, title, period phrase,
footnotes. Values are the cell values: numbers stay numbers, a published dash is a dash (value None, flag "dash"), text
that is not a number stays text (value None, flag "text"); nothing is repaired or computed. Layout surprises raise or are
reported in `problems`; the caller decides what to keep.
"""
import io
import re
import unicodedata
from calendar import monthrange
from datetime import date

import openpyxl

SUPER = "¹²³⁴⁵⁶⁷⁸⁹⁰⁺*†‡§"
DASHES = {"-", "--", "—", "–", "−", "‐"}
NA_TEXT = {"N/A", "NA", "(NA)", "n/a", "NR", "(NR)"}
MONTHS = {m: i for i, m in enumerate(["January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"], 1)}
# "Table C-4A.", "BAPCPA Table 1A.", "Wiretap Table 2—", "Table 1.1" (Judicial Facts and Figures)
TABLE_ID = re.compile(
    r"^\s*(?:(BAPCPA|Wiretap)\s+)?(?:Table|Report)\s+([A-Z]{1,4}(?:-?\d+[A-Z]?(?:-\d+)?)?(?:-[A-Za-z]+)*|\d+(?:\.\d+)?[A-Z]?)\s*[.\-:–—]?",
    re.I,
)
PREFIX_ID = {"BAPCPA": "BAPCPA", "WIRETAP": "WIRE"}


def printed_table_id(match):
    """The table id a TABLE_ID match names, upper case, with the publication prefix the listing uses ("BAPCPA 1A", "WIRE 2")."""
    prefix = match.group(1)
    ident = (match.group(2) or "").upper()
    return f"{PREFIX_ID[prefix.upper()]} {ident}" if prefix else ident


PERIOD_PHRASE = re.compile(
    r"((?:During|For)\s+the\s+(\d+|[A-Za-z]+)[-\s]Months?\s+Periods?\s+Ending|(?:During|For)\s+the\s+(?:\d+|[A-Za-z]+)[-\s]Month\s+Period\s+Ending|Quarter\s+Ended|Period\s+Ending|As\s+of|Ended|Ending)\s+"
    r"([A-Z][a-z]+\.?)\s+(\d{1,2}),?\s+(\d{4})"
    r"(?:\s*,?\s*(?:and|through|to)\s+(?:[A-Z][a-z]+\.?\s+\d{1,2},?\s+)?(\d{4}))?",
    re.I,
)


_SUP_TO_PUA = {ch: chr(0xE000 + i) for i, ch in enumerate("¹²³⁴⁵⁶⁷⁸⁹⁰")}
_PUA_TO_SUP = {v: k for k, v in _SUP_TO_PUA.items()}


def norm_text(value):
    """Cell text with the export's carriage-return escapes removed and whitespace normalised (newlines kept as spaces).
    Superscript digits (footnote marks) survive the Unicode normalisation unchanged."""
    if value is None:
        return ""
    text = "".join(_SUP_TO_PUA.get(ch, ch) for ch in str(value))
    text = unicodedata.normalize("NFKC", text)
    text = "".join(_PUA_TO_SUP.get(ch, ch) for ch in text).replace("_x000D_", "").replace("\r", "")
    text = re.sub(r"[ \t ]+", " ", text)
    return re.sub(r"\s*\n\s*", " ", text).strip()


def strip_marks(text):
    """Remove footnote marks (superscript digits and symbols) from a label or heading; plain digits stay."""
    text = "".join(ch for ch in text if ch not in SUPER)
    return re.sub(r"\s+", " ", text).strip()


def slug(text):
    text = strip_marks(norm_text(text)).lower()
    text = re.sub(r"[^a-z0-9%]+", "_", text).strip("_")
    return text[:60] or "x"


SUPERDIGITS = "¹²³⁴⁵⁶⁷⁸⁹⁰"
MARK = re.compile("[" + SUPERDIGITS + "*†‡§]+")


HELPER_HEADINGS = {"sort order"}

# A row that opens the footnote block under the data: "NOTE: ...", "Source: ..." or a line that starts with a footnote mark.
NOTES_START = re.compile(r"^(?:[" + SUPERDIGITS + r"†‡§]|\*\s|Notes?\b|Source\b)", re.I)


def marks_in(value):
    """Footnote marks (runs of superscript digits) printed in a cell, in order, without repeats."""
    out = []
    for m in MARK.findall(str(value)) if value is not None else []:
        if m not in out:
            out.append(m)
    return out


def indent_of(raw):
    if raw is None:
        return 0
    text = str(raw).replace(" ", " ").replace("_x000D_", "")
    return len(text) - len(text.lstrip(" "))


def to_value(cell):
    """(value, flag) for one data cell: number -> (number, None); dash -> (None, 'dash'); text -> (None, 'text:<raw>')."""
    if cell is None:
        return None, "blank"
    if isinstance(cell, bool):
        return None, "text:" + str(cell)
    if isinstance(cell, (int, float)):
        return cell, None
    text = norm_text(cell)
    if text == "":
        return None, "blank"
    if text in DASHES:
        return None, "dash"
    plain = text.replace(",", "").replace("$", "")
    plain = "".join(ch for ch in plain if ch not in SUPER).strip()
    if re.fullmatch(r"-?\d+(\.\d+)?", plain):
        number = float(plain)
        return (int(number) if re.fullmatch(r"-?\d+", plain) else number), "textnum"
    if re.fullmatch(r"-?\d+(\.\d+)?%", plain):
        return float(plain[:-1]), "textpct"
    if text in NA_TEXT:
        return None, "na"
    return None, "text:" + text[:40]


def is_numberlike(cell):
    value, flag = to_value(cell)
    return value is not None or flag == "dash"


PERIOD_NO_DAY = re.compile(
    r"((?:During|For)\s+the\s+(?:\d+|[A-Za-z]+)[-\s]Months?\s+Periods?\s+Ending|Period\s+Ending|Quarter\s+Ended)\s+([A-Z][a-z]+)\.?,?\s+((?:19|20)\d{2})\b",
    re.I,
)


CALENDAR_YEAR = re.compile(r"(?:During|For|In)\s+(?:the\s+)?Calendar\s+Year\s+((?:19|20)\d{2})\b", re.I)


def _month_number(name):
    name = name.rstrip(".")
    return MONTHS.get(name) or next((n for m, n in MONTHS.items() if m.lower().startswith(name[:3].lower())), None)


def parse_period_phrase(text):
    """(phrase, 'YYYY-MM-DD') from a title such as 'During the 12-Month Period Ending September 30, 2025'.

    A period phrase ("... Period Ending <date>") wins over a later "as of <date>" (a data-extraction date some editions print
    after it); "as of <date>" is the period itself only for point-in-time tables (cases pending as of a date). A phrase with a
    month and year but no day ("Period Ending September, 2020") ends on that month's last day."""
    flat = re.sub(r"\s+", " ", text)
    cal = CALENDAR_YEAR.search(flat)
    if cal and not PERIOD_PHRASE.search(flat):
        return cal.group(0), f"{cal.group(1)}-12-31"
    matches = list(PERIOD_PHRASE.finditer(flat))
    chosen = next((m for m in matches if not m.group(1).lower().startswith("as")), None)
    if chosen is None:
        nd = PERIOD_NO_DAY.search(flat)
        if nd:
            month_num = _month_number(nd.group(2))
            if month_num:
                last = monthrange(int(nd.group(3)), month_num)[1]
                return nd.group(0), date(int(nd.group(3)), month_num, last).isoformat()
        chosen = matches[0] if matches else None
    if chosen is None:
        return None, None
    month_num = _month_number(chosen.group(3))
    if not month_num:
        return chosen.group(0), None
    try:
        # a range ("March 31, 2025 and 2026", "Through 2025") names the editions it spans: the period end is the last one
        end = date(int(chosen.group(6) or chosen.group(5)), month_num, int(chosen.group(4))).isoformat()
    except ValueError:
        end = None
    return chosen.group(0), end


def fill_merged(ws, grid, limit_row):
    """Copy a merged range's value across its span, only for rows above limit_row (the header region)."""
    for merged in ws.merged_cells.ranges:
        if merged.min_row > limit_row:
            continue
        value = grid[merged.min_row - 1][merged.min_col - 1] if merged.min_row - 1 < len(grid) else None
        if value is None:
            continue
        for r in range(merged.min_row, min(merged.max_row, limit_row) + 1):
            for c in range(merged.min_col, merged.max_col + 1):
                if r - 1 < len(grid) and c - 1 < len(grid[r - 1]) and grid[r - 1][c - 1] is None:
                    grid[r - 1][c - 1] = value


def first_data_row(grid, label_cols_guess=2):
    """Index of the first data row: a labelled row with >= 2 number-like cells where the label is not a year header."""
    for i, row in enumerate(grid):
        cells = list(row)
        label_idx = next((k for k, v in enumerate(cells[:label_cols_guess]) if v is not None and norm_text(v)), None)
        if label_idx is None:
            continue
        label = norm_text(cells[label_idx])
        if TABLE_ID.match(label) or label.lower().startswith("note"):
            continue
        rest = cells[label_idx + 1:]
        numbers = [c for c in rest if is_numberlike(c)]
        if len(numbers) < 2:
            continue
        # a row of 4-digit years is a heading, not data
        values = [to_value(c)[0] for c in rest if to_value(c)[0] is not None]
        if values and all(isinstance(v, (int, float)) and float(v).is_integer() and 1900 <= v <= 2100 for v in values) and len(values) == len(numbers):
            continue
        return i
    return None


def parse_sheet(ws, problems):
    grid = [list(r) for r in ws.iter_rows(values_only=True)]
    if not grid:
        return None
    width = max(len(r) for r in grid)
    for r in grid:
        r.extend([None] * (width - len(r)))
    formats = [[c.number_format for c in row] for row in ws.iter_rows()]
    start = first_data_row(grid)
    if start is None:
        problems.append(f"{ws.title}: no data rows found")
        return None
    fill_merged(ws, grid, start)
    # ---- title block: leading rows up to the first row that has a header-like second cell, joined
    title_lines, table_id, header_start = [], None, 0
    for i in range(start):
        row = grid[i]
        texts = [norm_text(v) for v in row if v is not None and norm_text(v)]
        if not texts:
            continue
        first = texts[0]
        distinct = []
        for t in texts:
            if t not in distinct:
                distinct.append(t)
        m = TABLE_ID.match(first)
        looks_title = len(distinct) == 1 and (m or title_lines or re.search(r"Period Ending|Cases|Courts|Summary|Filed|Filings|Statistics", first))
        if not title_lines and not m and len(distinct) > 1:
            header_start = i
            break
        # a title line starts in the first column; a merged column heading that happens to repeat one text starts further right
        if looks_title and (m or title_lines) and row[0] is not None and norm_text(row[0]):
            title_lines.append(first)
            if m and table_id is None:
                table_id = printed_table_id(m)
            header_start = i + 1
            continue
        header_start = i
        break
    title_raw = " ".join(title_lines)
    # titles may keep the table id in the first line only
    # ---- header rows
    header_rows = [grid[i] for i in range(header_start, start) if any(v is not None and norm_text(v) for v in grid[i])]
    data_rows = []
    note_lines = []
    last_data = start
    value_cols = set()
    for i in range(start, len(grid)):
        row = grid[i]
        label_candidates = [(k, v) for k, v in enumerate(row[:3]) if v is not None and norm_text(v)]
        nums = [k for k, v in enumerate(row) if is_numberlike(v) and k >= 1]
        if label_candidates and nums:
            data_rows.append(i)
            last_data = i
            value_cols.update(nums)
        elif label_candidates and data_rows and NOTES_START.match(norm_text(label_candidates[0][1])):
            # the footnote block starts here; labelled rows with numbers below it ("Chapter 9: MO, E 1 85") belong to the notes
            break
    # ---- label columns: those before the first value column
    first_value = min(value_cols) if value_cols else 1
    label_cols = list(range(0, max(first_value, 1)))
    value_cols = sorted(c for c in value_cols if c >= first_value)
    # a "Sort Order" column is the publisher's helper for ordering rows (text codes such as "10H"), not a published statistic
    value_cols = [
        c for c in value_cols
        if not ([strip_marks(norm_text(hr[c])) for hr in header_rows if strip_marks(norm_text(hr[c]))][-1:] or [""])[0].lower() in HELPER_HEADINGS
    ]
    # ---- column paths
    columns = []
    for c in value_cols:
        path = []
        for hr in header_rows:
            t = strip_marks(norm_text(hr[c]))
            if t and (not path or path[-1] != t):
                path.append(t)
        fmts = [formats[i][c] for i in data_rows if c < len(formats[i]) and grid[i][c] is not None and not isinstance(grid[i][c], str)]
        fmt = max(set(fmts), key=fmts.count) if fmts else None
        col_marks = []
        for hr in header_rows:
            for m in marks_in(hr[c]):
                if m not in col_marks:
                    col_marks.append(m)
        columns.append({"index": c, "path": path, "fmt": fmt, "marks": col_marks})
    # ---- rows
    rows = []
    section = None
    for i in range(start, last_data + 1):
        row = grid[i]
        raw = next((v for k, v in enumerate(row[:max(len(label_cols), 1)]) if v is not None and norm_text(v)), None)
        has_values = any(is_numberlike(row[c]) for c in value_cols)
        if raw is None:
            continue
        label = strip_marks(norm_text(raw))
        if not has_values:
            # a heading row inside the table ("Contract Actions"): context for the rows that follow
            rows.append({"label_raw": str(raw), "label": label, "indent": indent_of(raw), "section": True, "values": None, "flags": None, "row": i + 1, "marks": marks_in(raw)})
            continue
        values, flags, cell_marks = [], [], {}
        for pos, c in enumerate(value_cols):
            v, f = to_value(row[c])
            values.append(v)
            flags.append(f)
            if isinstance(row[c], str) and marks_in(row[c]):
                cell_marks[str(pos)] = marks_in(row[c])
        rows.append({"label_raw": str(raw), "label": label, "indent": indent_of(raw), "section": False, "values": values, "flags": flags, "row": i + 1,
                     "marks": marks_in(raw), "cell_marks": cell_marks})
    # ---- notes after the data: raw notes keep their own line breaks, re-read from the cells
    raw_notes = []
    for i in range(last_data + 1, len(grid)):
        cells = [v for v in grid[i] if v is not None and str(v).strip()]
        if len(cells) > 1 and any(not isinstance(v, str) or re.fullmatch(r"[\d.,]+", v.strip()) for v in cells):
            # a small table printed under the notes ("Chapter 9: MO, E 1 ..."): keep each row as one line, in cell order
            line = "  ".join(norm_text(v) for v in cells)
            if line not in raw_notes:
                raw_notes.append(line)
            continue
        for v in grid[i]:
            if v is not None and str(v).strip():
                parts = [norm_text(p) for p in str(v).replace("_x000D_", "").replace("\r", "").split("\n")]
                for p in parts:
                    if p and p not in raw_notes:
                        raw_notes.append(p)
    note_marks = {}
    for line in raw_notes:
        m = re.match("^(" + MARK.pattern + r")\s*(.*)$", line)
        if m and m.group(2):
            note_marks.setdefault(m.group(1), line)
    # the title often carries the period; look in title lines and the first header cells
    phrase, period_end = parse_period_phrase(title_raw)
    if not phrase:
        for hr in header_rows[:2]:
            blob = " ".join(norm_text(v) for v in hr if v)
            phrase, period_end = parse_period_phrase(blob)
            if phrase:
                break
    return {
        "sheet": ws.title,
        "table_id": table_id,
        "title_raw": title_raw,
        "period_phrase": phrase,
        "period_end_from_title": period_end,
        "label_header": [strip_marks(norm_text(hr[0])) for hr in header_rows if norm_text(hr[0])],
        "columns": columns,
        "rows": rows,
        "notes": raw_notes,
        "note_marks": note_marks,
        "header_rows": len(header_rows),
    }


def parse_workbook(data):
    """Parse every sheet of an .xlsx workbook; returns (sheets, problems)."""
    problems = []
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    sheets = []
    for ws in wb.worksheets:
        parsed = parse_sheet(ws, problems)
        if parsed:
            sheets.append(parsed)
    return sheets, problems
