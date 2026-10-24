"""Offline tests for court ids the tables print (exact matches only) and the combined Arkansas bankruptcy row."""
import copy
import io
import unittest

import openpyxl

import build_tables as B
import build_uscourts_tables as BU
import normalize as N
import test_normalize as T
import utx_parse as P


class CourtIdTests(unittest.TestCase):
    def test_codes_that_name_no_court_are_never_given_an_id(self):
        # Judicial Business S-19 FY2017 prints AK,E and AK,W in the Eighth Circuit block and AR in the Ninth
        self.assertEqual(N.classify_row("district", "AK,E"), ("district", "unmatched:AK,E"))
        self.assertEqual(N.classify_row("district", "AK,W"), ("district", "unmatched:AK,W"))
        self.assertEqual(N.classify_row("district", "AR"), ("district", "unmatched:AR"))
        self.assertEqual(N.classify_row("district", "AK"), ("district", "akd"))
        self.assertEqual(N.classify_row("district", "AR,E"), ("district", "ared"))

    def test_the_former_eastern_district_of_illinois(self):
        # "IL,E 1" with footnote 1 "No longer a separate district": the directory's District Court, E.D. Illinois
        self.assertEqual(N.classify_row("district", "IL,E 1"), ("district", "illinoised"))
        self.assertIn("illinoised", BU.TITLES["district_titles"])
        self.assertNotIn("illinoised", BU.TITLES["districts"])  # not one of the 94 current districts

    def test_arkansas_in_a_bankruptcy_table_is_the_combined_row(self):
        self.assertEqual(N.classify_row("bankruptcy", "AR ²"), ("district", "arcomb"))
        self.assertEqual(N.classify_row("bankruptcy", "AR, E"), ("district", "ared"))
        # never the Arizona bankruptcy court (directory id arb) or the historical District of Arkansas (ard)
        self.assertNotIn(N.classify_row("district", "AR")[1], ("arb", "ard"))

    def test_unmatched_rows_stay_in_their_circuit_block_for_the_rollup(self):
        n = copy.deepcopy(T.normalized())
        for r in n["rows"]:
            if r["entity_id"] == "ared":
                r["entity_id"] = "unmatched:AK,E"
        import checks as C
        self.assertEqual(C.rollup_failures(n), [])


def stray_note_workbook():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "pdf for bk iling"
    ws["A2"] = ("\nTable F-2 Monthly.\nU.S. Bankruptcy Courts—Business and Nonbusiness Cases Commenced, by Chapter of the Bankruptcy Code,\n"
                "During the One-Month Period Ending September 30, 2021, Based on Data Current as of September 30, 2021")
    ws.merge_cells("A2:D2")
    ws["A3"], ws["B3"], ws["C3"], ws["D3"] = "Circuit and District", "Total All Chapters", "Total Chapter 7", "Total Chapter 13"
    ws.append(["Total", 30, 20, 10])
    ws.append(["DC", 3, 2, 1])
    ws.append(["AR ²", 27, 18, 9])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class StrayNoteTests(unittest.TestCase):
    def test_a_note_above_the_title_is_not_the_title_or_a_heading(self):
        sheets, _ = P.parse_workbook(stray_note_workbook())
        s = sheets[0]
        self.assertTrue(s["title_raw"].startswith("Table F-2 Monthly."))
        self.assertEqual(s["period_end_from_title"], "2021-09-30")
        self.assertEqual(s["stray_text"], ["pdf for bk iling"])
        self.assertEqual([c["path"] for c in s["columns"]], [["Total All Chapters"], ["Total Chapter 7"], ["Total Chapter 13"]])


class ArkansasTests(unittest.TestCase):
    def docs(self, layout):
        """layout per period: 'split' prints AR,E and AR,W, 'combined' prints AR."""
        out = []
        for i, kind in enumerate(layout):
            rows = [{"entity": "national", "entity_id": "national", "section": False, "values": [1]}]
            if kind == "split":
                rows += [{"entity": "district", "entity_id": "ared", "section": False, "values": [1]},
                         {"entity": "district", "entity_id": "arwd", "section": False, "values": [1]}]
            else:
                rows += [{"entity": "district", "entity_id": "arcomb", "section": False, "values": [2]}]
            out.append({"norm": {"rows": rows, "period_end": f"201{i}-09-30", "columns": []}})
        for d in out:
            d["rows_by_entity"], _ = B.row_index(d["norm"])
        return out

    def test_the_district_records_say_where_the_combined_row_begins(self):
        docs = self.docs(["split", "split", "combined", "combined"])
        got = BU.arkansas_combined("ared", docs)
        self.assertEqual(got["reportedWith"]["court"], "arcomb")
        self.assertEqual(got["reportedWith"]["from"], "2012-09-30")
        self.assertEqual(got["reportedWith"]["lastOwn"], "2011-09-30")
        self.assertEqual(BU.arkansas_combined("arcomb", docs)["covers"], ["ared", "arwd"])
        self.assertEqual(BU.arkansas_combined("nysd", docs), {})

    def test_no_combined_row_no_note(self):
        self.assertEqual(BU.arkansas_combined("ared", self.docs(["split", "split"])), {})


if __name__ == "__main__":
    unittest.main()
